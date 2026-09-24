#!/usr/bin/env python
"""Execute the legacy PIT generator in a temporary copy and compare key results.

The legacy UI is never invoked.  Passing --company-id and --period-id additionally
compares the current database-backed engine result directly, without HTTP.
"""
from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl


def money(value):
    if value is None or value == "": return None
    try: return Decimal(str(value).replace(",", "")).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError): return None


def text(value): return "" if value is None else str(value).strip()
def money_equal(left, right, tolerance=Decimal("0.01")):
    return left is None and right is None or left is not None and right is not None and abs(left - right) <= tolerance


def find_sheet(workbook, token):
    return next((sheet for sheet in workbook.worksheets if token in sheet.title), None)


def extract_tax_amount(sheet):
    if sheet is None: return {}
    result = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        org, subject = text(row[0] if len(row) > 0 else None), text(row[2] if len(row) > 2 else None)
        if not org or subject not in {"21510006", "21510008", "21510009", "21510016"}: continue
        result[(org, subject)] = {"declared_tax_amount": money(row[9] if len(row) > 9 else None), "credit_amount": money(row[7] if len(row) > 7 else None), "current_difference": money(row[10] if len(row) > 10 else None), "closing_balance": money(row[8] if len(row) > 8 else None), "cumulative_difference": money(row[12] if len(row) > 12 else None), "business_declared_difference": money(row[15] if len(row) > 15 else None)}
    return result


def run_legacy(fixture_dir: Path):
    candidates = list(fixture_dir.glob("tax_reconciliation*.py"))
    if not candidates: raise RuntimeError("未找到 legacy tax_reconciliation.py")
    spec = importlib.util.spec_from_file_location("legacy_tax_reconciliation", candidates[0]); module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None; spec.loader.exec_module(module)
    template = next((path for path in fixture_dir.glob("*核对底稿*.xlsx") if "结果" not in path.name), None)
    org = next(iter(fixture_dir.glob("*机构信息表*.xlsx")), None); salaries = sorted(path for path in fixture_dir.glob("*工资单*.xlsx") if "经纪人" not in path.name); broker = next(iter(fixture_dir.glob("*经纪人*.xlsx")), None)
    if not template or not org or len(salaries) < 2 or not broker: raise RuntimeError("legacy fixture 缺少模板、机构信息表或三类工资表")
    five = next((path for path in salaries if "5位" in path.name), salaries[0]); seven = next((path for path in salaries if "7位" in path.name), salaries[1])
    with TemporaryDirectory(prefix="pit_legacy_") as temporary:
        root = Path(temporary); target = root / template.name; shutil.copy2(template, target)
        loader = module.DataLoader()
        loader.load_org_info(org)
        source_workbook = openpyxl.load_workbook(target, data_only=True, read_only=True)
        try:
            loader.load_source_from_template(source_workbook)
        finally:
            source_workbook.close()
        loader.load_payroll(five, seven, broker)
        generator = module.ReconciliationGenerator(target, loader); generator.generate_all(mode="pre"); output = root / "legacy_result.xlsx"; generator.save(output)
        workbook = openpyxl.load_workbook(output, data_only=True, read_only=True)
        try:
            return {"tax_amount": extract_tax_amount(find_sheet(workbook, "个税明细税额核对")), "sheets": [sheet.title for sheet in workbook.worksheets]}
        finally:
            workbook.close()


def production_tax_amount(company_id: int, period_id: int):
    from app.core.company_context import reset_company_id, set_company_id
    from app.db.session import SessionLocal
    from app.services.pit_reconciliation.engine import PitReconciliationEngine
    from app.services.pit_reconciliation.source_service import PitSourceService
    token = set_company_id(company_id); db = SessionLocal()
    try:
        result = PitReconciliationEngine().calculate(PitSourceService(db, company_id, period_id).load_bundle())
        return {(row["org_code"], row["subject_code"]): row for row in result["tax_checks"]}
    finally:
        db.close(); reset_company_id(token)


def markdown(rows, compared):
    lines = ["# PIT Legacy Regression", "", f"Overall: {'PASS' if compared and all(item[-1] == 'PASS' for item in rows) else 'NOT COMPARED' if not compared else 'FAIL'}", "", "## 个税明细税额核对", "", "| Key | Field | Legacy | New | Diff | Result |", "| --- | --- | ---: | ---: | ---: | --- |"]
    for key, field, legacy, current, result in rows:
        diff = "—" if legacy is None or current is None else f"{current - legacy:.2f}"
        lines.append(f"| {key[0]} / {key[1]} | {field} | {legacy if legacy is not None else 'NULL'} | {current if current is not None else 'NULL'} | {diff} | {result} |")
    return "\n".join(lines) + "\n"


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--fixture-dir", required=True); parser.add_argument("--company-id", type=int); parser.add_argument("--period-id", type=int); parser.add_argument("--output-dir", default="outputs")
    args = parser.parse_args(); fixture = Path(args.fixture_dir); legacy = run_legacy(fixture); compared = args.company_id is not None and args.period_id is not None; rows = []
    current = production_tax_amount(args.company_id, args.period_id) if compared else {}
    for key, legacy_values in sorted(legacy["tax_amount"].items()):
        for field in ("declared_tax_amount", "credit_amount", "current_difference", "closing_balance", "cumulative_difference", "business_declared_difference"):
            old = legacy_values[field]; new = money(current.get(key, {}).get(field)) if compared else None; rows.append((key, field, old, new, "PASS" if compared and money_equal(old, new) else "NOT COMPARED" if not compared else "FAIL"))
    output = Path(args.output_dir); output.mkdir(parents=True, exist_ok=True); report = {"legacy_mode": "execute", "legacy_sheets": legacy["sheets"], "compared": compared, "company_id": args.company_id, "period_id": args.period_id, "rows": [{"key": list(key), "field": field, "legacy": str(old) if old is not None else None, "new": str(new) if new is not None else None, "result": result} for key, field, old, new, result in rows]}
    (output / "pit_legacy_regression_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"); (output / "pit_legacy_regression_report.md").write_text(markdown(rows, compared), encoding="utf-8")
    print(json.dumps({"legacy_mode": "execute", "compared": compared, "rows": len(rows), "output": str(output)}, ensure_ascii=False))


if __name__ == "__main__": main()
