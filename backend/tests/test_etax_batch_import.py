from __future__ import annotations

import importlib
import sys
import types
from pathlib import Path

from openpyxl import Workbook, load_workbook


def _module():
    if "playwright.sync_api" not in sys.modules:
        playwright = types.ModuleType("playwright")
        sync_api = types.ModuleType("playwright.sync_api")
        sync_api.Page = type("Page", (), {})
        sync_api.TimeoutError = TimeoutError
        sync_api.sync_playwright = lambda: None
        playwright.sync_api = sync_api
        sys.modules["playwright"] = playwright
        sys.modules["playwright.sync_api"] = sync_api
    vendor = Path(__file__).parents[1] / "vendor" / "etax_rpa"
    path_text = str(vendor)
    if path_text not in sys.path:
        sys.path.insert(0, path_text)
    return importlib.import_module("etax_batch_import")


def test_build_tasks_separates_intern_declaration_from_normal_salary(tmp_path):
    module = _module()
    tasks = module.build_tasks("10001")
    normal = next(task for task in tasks if task.label.endswith("正常工资薪金所得"))
    intern = next(task for task in tasks if "劳务报酬" in task.label and task.label.endswith("实习生）"))

    intern_file = tmp_path / "10001_个税申报表_实习生(2026年08月).xlsx"
    intern_file.write_bytes(b"placeholder")

    assert module.find_one_file(tmp_path, normal.patterns, normal.exclude_patterns) is None
    assert module.find_one_file(tmp_path, intern.patterns, intern.exclude_patterns) == intern_file.resolve()
    assert next(task for task in tasks if task.label == "人员信息采集-员工").clear_before_import is False
    assert next(task for task in tasks if task.label == "人员信息采集-实习生").clear_before_import is False
    assert next(task for task in tasks if task.label == "综合所得申报-劳务报酬（实习生）").clear_before_import is False
    assert next(task for task in tasks if task.label == "综合所得申报-正常工资薪金所得").clear_before_import is True
    assert next(task for task in tasks if task.label == "限售股所得申报").clear_before_import is False


def test_clear_data_waits_cover_delayed_tax_site_controls():
    module = _module()

    assert module.ACTION_BAR_WAIT_SECONDS >= 30
    assert module.DROPDOWN_MENU_WAIT_SECONDS >= 30
    assert module.CLEAR_CONFIRMATION_WAIT_SECONDS >= 30


def test_read_orgs_preserves_search_result_index(tmp_path):
    module = _module()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["机构名称", "机构代码", "RPA搜索结果序号"])
    sheet.append(["重复机构", "10001", 2])
    path = tmp_path / "机构.xlsx"
    workbook.save(path)

    orgs = module.read_orgs_from_excel(path)

    assert [(org.name, org.code, org.search_result_index) for org in orgs] == [("重复机构", "10001", 2)]


def test_summary_cells_follow_headers_when_tax_site_changes_column_order():
    module = _module()
    headers = ["收入合计（元）", "所得项目", "应补/退税额（元）", "填报人次"]
    cells = ["1,000.00", "工资薪金所得", "100.00", "2"]

    assert module.map_summary_cells(headers, cells) == {
        "所得项目": "工资薪金所得",
        "填写人次": "2",
        "收入合计（元）": "1,000.00",
        "应补/退税额（元）": "100.00",
    }


def test_task_progress_survives_failed_organization_and_is_skipped_on_resume(tmp_path, monkeypatch):
    module = _module()
    monkeypatch.setattr(module, "OUTPUT_DIR", tmp_path)
    org = module.TaxOrg(code="10001", name="测试机构")
    progress = {"month": "2026-08", "completed": {}, "tasks": {}}
    task = module.build_tasks(org.code)[0]

    module.mark_task_completed("2026-08", org, task, progress)
    assert module.completed_task_ids(progress, org) == {task.task_id}
    assert "10001" not in progress["completed"]

    module.reset_progress_for_orgs("2026-08", [org])
    assert module.completed_task_ids(module.load_progress("2026-08"), org) == set()


def test_reconciliation_workbook_preserves_other_org_sheets_and_replaces_current(tmp_path, monkeypatch):
    module = _module()
    monkeypatch.setattr(module, "OUTPUT_DIR", tmp_path)
    org_a = module.TaxOrg(name="机构A", code="10001")
    org_b = module.TaxOrg(name="机构B", code="10002")

    module.write_declaration_check_excel("2026-08", org_a, [("综合所得", [{"所得项目": "工资", "填写人次": "1", "收入合计（元）": "100", "应补/退税额（元）": "10"}])])
    module.write_declaration_check_excel("2026-08", org_b, [("综合所得", [{"所得项目": "工资", "填写人次": "2", "收入合计（元）": "200", "应补/退税额（元）": "20"}])])
    module.write_declaration_check_excel("2026-08", org_a, [("综合所得", [{"所得项目": "工资", "填写人次": "3", "收入合计（元）": "300", "应补/退税额（元）": "30"}])])

    workbook = load_workbook(tmp_path / "2026-08个税申报核对.xlsx", data_only=True)
    assert workbook.sheetnames == ["10002_机构B", "10001_机构A"]
    assert workbook["10002_机构B"]["C2"].value == "2"
    assert workbook["10001_机构A"]["C2"].value == "3"
    assert not list(tmp_path.glob(".*.tmp-*.xlsx"))
    workbook.close()


def test_process_org_clears_each_scope_once_and_records_each_upload(tmp_path, monkeypatch):
    module = _module()
    org = module.TaxOrg(name="机构A", code="10001")
    first = module.ImportTask("综合所得申报", "劳务报酬", ["first.xlsx"], "任务一")
    second = module.ImportTask("综合所得申报", "劳务报酬", ["second.xlsx"], "任务二")
    (tmp_path / "first.xlsx").write_bytes(b"1")
    (tmp_path / "second.xlsx").write_bytes(b"2")
    progress = {"month": "2026-08", "completed": {}, "tasks": {}}
    cleared = []
    uploaded = []
    completed = []

    monkeypatch.setattr(module, "build_tasks", lambda _code: [first, second])
    monkeypatch.setattr(module, "input_dirs_for_org", lambda *_args: [tmp_path])
    monkeypatch.setattr(module, "ensure_withholding_page", lambda page: page)
    monkeypatch.setattr(module, "switch_org", lambda *_args: None)
    monkeypatch.setattr(module, "enter_declaration_page", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(module, "clear_existing_data", lambda _page, task: cleared.append(task.label))
    monkeypatch.setattr(module, "upload_and_verify", lambda _page, path, **_kwargs: uploaded.append(path.name))
    monkeypatch.setattr(module, "mark_task_completed", lambda _month, _org, task, _progress: completed.append(task.label))
    monkeypatch.setattr(module, "run_post_import_check", lambda *_args: None)

    module.process_org(object(), org, "2026-08", tmp_path, 30, progress=progress)

    assert cleared == ["任务一"]
    assert uploaded == ["first.xlsx", "second.xlsx"]
    assert completed == ["任务一", "任务二"]
