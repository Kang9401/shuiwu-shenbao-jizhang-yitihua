from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.db.session import Base
from app.models.core import Period, UploadedFile
from app.models.accounting import PersonnelMasterArtifact
from app.services.personnel_update import PERSONNEL_COLLECTION_COLUMNS
from app.workflows import get_workflow, list_workflows


def _db_session():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    testing_session_local = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    Base.metadata.create_all(bind=engine)
    return testing_session_local()


def test_all_planned_workflows_registered():
    codes = {workflow.info.code for workflow in list_workflows()}
    assert {
        "general_salary_tax",
        "staff_info_update",
        "annual_bonus_tax",
        "restricted_stock_interest_tax",
        "intern_tax",
        "broker_tax",
        "invoice_booking",
        "vat_deduction",
        "voucher_draft",
    }.issubset(codes)


def test_unknown_workflow_raises_clear_error():
    try:
        get_workflow("missing")
    except ValueError as exc:
        assert "未知流程" in str(exc) or "鏈煡娴佺▼" in str(exc)
    else:
        raise AssertionError("expected ValueError")


def test_general_salary_tax_exposes_single_salary_sheet_role():
    workflow = get_workflow("general_salary_tax")
    assert workflow.info.required_file_roles == ["salary_sheet"]
    assert not {
        "marketing_salary_sheet",
        "rank_salary_sheet",
        "branch_salary_sheet",
        "headquarters_salary_sheet",
        "digital_ops_salary_sheet",
        "advisor_salary_sheet",
    }.intersection(workflow.info.required_file_roles)


def test_registered_tax_workflow_writes_real_artifacts(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_root", tmp_path / "storage")
    source = tmp_path / "salary.xlsx"
    pd.DataFrame([{"employee": "1001"}]).to_excel(source, index=False)

    uploaded = UploadedFile(
        period_id=None,
        file_role="salary_sheet",
        original_name="salary.xlsx",
        stored_path=str(source),
        size_bytes=source.stat().st_size,
        validation_status="uploaded",
        validation_issues=[],
    )

    result = get_workflow("general_salary_tax").run(None, 9001, None, [uploaded])

    assert result.status == "needs_review"
    assert result.summary["artifacts"] >= 1
    assert all(Path(path).exists() for _, path in result.artifact_paths)


def test_tax_workflow_uses_shared_monthly_staff_info_when_not_uploaded(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_root", tmp_path / "storage")
    db = _db_session()
    period = Period(year=2025, month=2, name="2025-02")
    db.add(period)
    db.commit()
    db.refresh(period)

    tax_path = tmp_path / "bonus.xlsx"
    staff_path = tmp_path / "staff.xlsx"
    pd.DataFrame().to_excel(tax_path, index=False)
    pd.DataFrame([{"staff": "shared"}]).to_excel(staff_path, index=False)
    db.add(
        PersonnelMasterArtifact(
            period_id=period.id,
            person_type="employee",
            scope_type="month",
            scope_code="",
            file_name="staff.xlsx",
            stored_path=str(staff_path),
            row_count=1,
            validation_issues=[],
        )
    )
    db.commit()

    uploaded = UploadedFile(
        period_id=period.id,
        file_role="bonus_sheet",
        original_name="bonus.xlsx",
        stored_path=str(tax_path),
        size_bytes=tax_path.stat().st_size,
        validation_status="uploaded",
        validation_issues=[],
    )

    result = get_workflow("annual_bonus_tax").run(db, 9002, period.id, [uploaded])

    assert result.summary["shared_staff_info"] is True
    assert result.summary["input_files"] == 2
    assert all(Path(path).exists() for _, path in result.artifact_paths)
    db.close()


def test_broker_tax_copies_previous_month_master_before_using_shared_staff_info(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_root", tmp_path / "storage")
    db = _db_session()
    previous = Period(year=2025, month=1, name="2025-01")
    current = Period(year=2025, month=2, name="2025-02")
    db.add_all([previous, current])
    db.commit()
    db.refresh(previous)
    db.refresh(current)

    staff_path = tmp_path / "broker_staff.xlsx"
    income_path = tmp_path / "broker_income.xlsx"
    pd.DataFrame([{
        "员工编号": "1001", "姓名": "张三", "证件类型": "居民身份证",
        "证件号码": "110101199001010011", "分支机构代码": "10301",
    }]).to_excel(staff_path, index=False)
    pd.DataFrame([{"员工编号": "1001", "本期收入": 100, "个人所得税": 10}]).to_excel(income_path, index=False)
    db.add(PersonnelMasterArtifact(
        period_id=previous.id, person_type="broker", scope_type="month", scope_code="",
        file_name=staff_path.name, stored_path=str(staff_path), row_count=1, validation_issues=[],
    ))
    db.commit()

    uploaded = UploadedFile(
        period_id=current.id, file_role="broker_income", original_name=income_path.name,
        stored_path=str(income_path), size_bytes=income_path.stat().st_size,
        validation_status="uploaded", validation_issues=[],
    )
    result = get_workflow("broker_tax").run(db, 9010, current.id, [uploaded])

    assert result.summary["shared_staff_info"] is True
    assert db.query(PersonnelMasterArtifact).filter(
        PersonnelMasterArtifact.period_id == current.id,
        PersonnelMasterArtifact.person_type == "broker",
    ).count() == 1
    declaration = next(Path(path) for kind, path in result.artifact_paths if kind == "declaration")
    columns = list(pd.read_excel(declaration).columns)
    assert columns[columns.index("减免税额"):columns.index("备注") + 1] == ["减免税额", "协定减免", "备注"]
    broker_book = load_workbook(declaration, data_only=True)
    assert broker_book["Sheet1"]["F2"].data_type == "s"
    broker_book.close()
    db.close()


def test_restricted_stock_generates_declaration_before_balance_reconciliation(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_root", tmp_path / "storage")
    db = _db_session()
    period = Period(year=2025, month=2, name="2025-02")
    db.add(period)
    db.commit()
    db.refresh(period)
    source = tmp_path / "restricted_stock.xlsx"
    pd.DataFrame([
        {
            "机构代码": "10301",
            "客户姓名": "张三",
            "证件号码": "110101199001010011",
            "证件类型": "居民身份证",
            "证券账户号": "A001",
            "证券代码": "000001",
            "证券名称": "示例股票",
            "成交价格": 10,
            "成交数量": 100,
            "原值总金额": 200,
        },
    ]).to_excel(source, index=False)
    uploaded = UploadedFile(
        period_id=period.id,
        file_role="tax_sheet",
        original_name=source.name,
        stored_path=str(source),
        size_bytes=source.stat().st_size,
        validation_status="uploaded",
        validation_issues=[],
    )

    result = get_workflow("restricted_stock_interest_tax").run(db, 9003, period.id, [uploaded])

    assert result.status == "success"
    assert any(kind == "declaration" for kind, _ in result.artifact_paths)
    collection_paths = [Path(path) for kind, path in result.artifact_paths if kind == "personnel_collection"]
    assert collection_paths
    assert "_2-人员信息采集_客户" in collection_paths[0].name
    assert result.summary["generated_files"] == 2
    assert result.summary["restricted_stock_declared_amount"] == 800
    assert result.summary["restricted_stock_withheld_tax"] == 160
    db.close()


def test_restricted_stock_and_interest_use_separate_tax_templates(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_root", tmp_path / "storage")
    db = _db_session()
    period = Period(year=2025, month=2, name="2025-02")
    db.add(period)
    db.commit()
    db.refresh(period)

    stock_path = tmp_path / "限售股申报.xlsx"
    interest_path = tmp_path / "债券兑息扣税.xlsx"
    pd.DataFrame([{
        "机构代码": "10301", "客户姓名": "张三", "证件类型": "身份证", "证件号码": "110101199001010011",
        "证券账户号": "A001", "证券代码": "000001", "证券名称": "示例股票", "企业名称": "示例公司",
        "成交价格": 10, "成交数量": 100, "原值总金额": 200,
    }]).to_excel(stock_path, index=False)
    pd.DataFrame([{
        "机构代码": "10301", "客户姓名": "李四", "证件类型": "居民身份证", "证件号码": "110101199001010022",
        "债券兑息": 100, "税率（%）": 20,
    }]).to_excel(interest_path, index=False)
    uploads = [
        UploadedFile(period_id=period.id, file_role="tax_sheet", original_name=stock_path.name, stored_path=str(stock_path), size_bytes=stock_path.stat().st_size, validation_status="uploaded", validation_issues=[]),
        UploadedFile(period_id=period.id, file_role="interest_tax", original_name=interest_path.name, stored_path=str(interest_path), size_bytes=interest_path.stat().st_size, validation_status="uploaded", validation_issues=[]),
    ]

    result = get_workflow("restricted_stock_interest_tax").run(db, 9004, period.id, uploads)
    declarations = [Path(path) for kind, path in result.artifact_paths if kind == "declaration"]
    assert {path.name for path in declarations} == {
        "10301_7-限售股转让所得(2025年02月).xlsx",
        "10301_6-利息、股息、红利所得(2025年02月).xlsx",
    }
    assert all(pd.ExcelFile(path).sheet_names == ["Sheet1"] for path in declarations)
    sheets = {path.name: pd.read_excel(path, sheet_name="Sheet1") for path in declarations}
    assert list(sheets["10301_7-限售股转让所得(2025年02月).xlsx"].columns) == [
        "工号", "*姓名", "*证件类型", "*证件号码", "*证券账户号", "*股票代码", "*股票名称",
        "上市公司纳税人识别号", "上市公司名称", "上市公司主管税务机关", "*每股计税价格(元/股)",
        "*转让股数(股)", "限售股原值", "合理税费", "准予扣除的捐赠额", "协定减免",
    ]
    assert list(sheets["10301_6-利息、股息、红利所得(2025年02月).xlsx"].columns) == [
        "工号", "*姓名", "*证件类型", "*证件号码", "*所得项目", "上市板块", "*收入", "免税收入",
        "准予扣除的捐赠额", "*税率", "协定税率", "减免税额", "协定减免", "备注",
    ]
    assert sheets["10301_6-利息、股息、红利所得(2025年02月).xlsx"].iloc[0]["*税率"] == "20%"
    assert sheets["10301_7-限售股转让所得(2025年02月).xlsx"].iloc[0]["*每股计税价格(元/股)"] == 10
    assert sheets["10301_6-利息、股息、红利所得(2025年02月).xlsx"].iloc[0]["*收入"] == 100
    stock_book = load_workbook(next(path for path in declarations if "_7-" in path.name), data_only=True)
    interest_book = load_workbook(next(path for path in declarations if "_6-" in path.name), data_only=True)
    assert stock_book["Sheet1"]["K2"].data_type == "s"
    assert stock_book["Sheet1"]["L2"].data_type == "s"
    assert interest_book["Sheet1"]["G2"].data_type == "s"
    stock_book.close()
    interest_book.close()
    collection = next(Path(path) for kind, path in result.artifact_paths if kind == "personnel_collection")
    collection_rows = pd.read_excel(collection, dtype=str).fillna("")
    stock_collection = collection_rows[collection_rows["*姓名"] == "张三"].iloc[0]
    interest_collection = collection_rows[collection_rows["*姓名"] == "李四"].iloc[0]
    assert stock_collection["*证件类型"] == "居民身份证"
    assert stock_collection["其他情况说明"] == "申报其他所得"
    assert interest_collection["其他情况说明"] == "扣缴申报利息股息红利所得"
    db.close()


def test_restricted_stock_does_not_write_tax_template_with_missing_required_fields(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_root", tmp_path / "storage")
    source = tmp_path / "restricted_stock.xlsx"
    pd.DataFrame([{
        "机构代码": "10301", "客户姓名": "张三", "证件号码": "110101199001010011",
        "证件类型": "居民身份证", "成交价格": 10, "成交数量": 100, "原值总金额": 200,
    }]).to_excel(source, index=False)
    uploaded = UploadedFile(
        period_id=None, file_role="tax_sheet", original_name=source.name, stored_path=str(source),
        size_bytes=source.stat().st_size, validation_status="uploaded", validation_issues=[],
    )

    result = get_workflow("restricted_stock_interest_tax").run(None, 9007, None, [uploaded])
    assert result.status == "needs_review"
    assert not any(kind == "declaration" for kind, _ in result.artifact_paths)
    assert result.summary["issues"] == 1


def test_restricted_stock_reconciliation_does_not_regenerate_declarations(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_root", tmp_path / "storage")
    source = tmp_path / "restricted_stock.xlsx"
    pd.DataFrame([{
        "机构代码": "10301", "客户姓名": "张三", "证件号码": "110101199001010011",
        "证件类型": "居民身份证", "成交价格": 10, "成交数量": 100, "原值总金额": 200,
    }]).to_excel(source, index=False)
    uploaded = UploadedFile(
        period_id=None,
        file_role="tax_sheet",
        original_name=source.name,
        stored_path=str(source),
        size_bytes=source.stat().st_size,
        validation_status="uploaded",
        validation_issues=[],
    )

    result = get_workflow("restricted_stock_interest_tax").run(None, 9005, None, [uploaded], operation="reconcile")
    assert result.status == "needs_review"
    assert not any(kind in {"declaration", "personnel_collection"} for kind, _ in result.artifact_paths)


def test_intern_workflow_writes_personnel_collection_and_declaration_templates(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_root", tmp_path / "storage")
    source = tmp_path / "实习生补贴.xlsx"
    pd.DataFrame([{
        "营业部名称": "广州昌岗中路", "*姓名": "张三", "*证件类型": "居民身份证",
        "*证件号码": "110101199001010011", "实习开始时间": "2025-02-01",
        "发放补贴数\n（元）": 500, "联系方式": "13800138000",
    }]).to_excel(source, index=False, startrow=1)
    uploaded = UploadedFile(
        period_id=None,
        file_role="intern_salary",
        original_name=source.name,
        stored_path=str(source),
        size_bytes=source.stat().st_size,
        validation_status="uploaded",
        validation_issues=[],
    )

    result = get_workflow("intern_tax").run(None, 9006, None, [uploaded])
    artifacts = {(kind, Path(path).name): Path(path) for kind, path in result.artifact_paths}
    declaration = next(path for (kind, _), path in artifacts.items() if kind == "declaration")
    collection = next(path for (kind, _), path in artifacts.items() if kind == "personnel_collection")
    assert "10367_4-个税申报表_实习生" in declaration.name
    assert "10367_2-人员信息采集_实习生" in collection.name
    collection_columns = [column for column in PERSONNEL_COLLECTION_COLUMNS if column != "人员状态"]
    collection_table = pd.read_excel(collection, dtype=str)
    assert list(collection_table.columns) == collection_columns
    assert not {"*所得项目", "本期收入", "实习生本期收入"}.intersection(collection_table.columns)
    intern_row = collection_table.loc[collection_table["*姓名"] == "张三"].iloc[0]
    assert intern_row["*任职受雇从业类型"] == "实习学生（全日制学历教育）"
    assert intern_row["任职受雇从业日期"] == "2025-02-01"
    intern_book = load_workbook(declaration, data_only=True)
    assert intern_book["Sheet1"]["F2"].data_type == "s"
    intern_book.close()
    assert list(pd.read_excel(declaration).columns) == [
        "工号", "*姓名", "*证件类型", "*证件号码", "*所得项目", "本期收入",
        "本期免税收入", "累计个人养老金", "商业健康保险", "税延养老保险", "其他",
        "允许扣除的税费", "减免税额", "协定减免", "备注",
    ]
