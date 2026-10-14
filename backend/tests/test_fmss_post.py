from decimal import Decimal
from types import SimpleNamespace

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.session import Base
from app.integrations.fmss.errors import FmssStateConflict, FmssImportFailed
from app.integrations.fmss.iit.schemas import FmssSheet
from app.models.accounting import ReconciliationImportBatch, ReconciliationImportRow
from app.models.core import Company, Period
from app.models.pit_reconciliation import PitBankTaxMatch, PitReconciliationOrgSummary, PitReconciliationWorkpaper
from app.core.config import settings
from app.services.pit_fmss_excel import PitFmssExcelBuilder
from app.services.pit_fmss_mapper import POST_SHEET_KEYS, PitFmssMapper
from app.services.pit_online_service import PitOnlineService


POST_HEADERS = {
    "iit_post_tax_check": ("机构代码", "营业部全称", "申报表", "完税证明", "申报表与完税证明差异金额11", "差异原因11", "银行流水个税", "完税证明与银行流水差异金额12", "差异原因12"),
    "iit_post_certificate": ("证明类型", "文件名", "纳税人名称", "纳税人识别号", "税种", "品目", "税款所属时期", "入(退)库日期", "实缴（退）金额"),
    "iit_post_bank_flow": tuple(f"列{index}" for index in range(29)),
    "iit_post_bank_tax": ("机构代码", "营业部全称", "本方账号", "交易时间", "交易摘要", "借方金额"),
}


def _sheets():
    return [FmssSheet(key=key, title=key, headers=POST_HEADERS[key]) for key in POST_SHEET_KEYS]


def _db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)()


def _post_context(tmp_path):
    db = _db()
    company = Company(name="成都分公司", code="17020", operator_name="测试")
    db.add(company)
    db.flush()
    period = Period(company_id=company.id, year=2026, month=7, name="2026-07")
    db.add(period)
    db.flush()
    workpaper = PitReconciliationWorkpaper(
        company_id=company.id, period_id=period.id, tax_type="pit", stage="post_payment",
        workflow_status="pending_submission", data_status="ready", calculation_status="success", rule_version="test",
    )
    db.add(workpaper)
    db.flush()
    pre = PitReconciliationWorkpaper(
        company_id=company.id, period_id=period.id, tax_type="pit", stage="pre_payment",
        workflow_status="reviewed", data_status="ready", calculation_status="success", rule_version="test", draft_revision=1,
    )
    db.add(pre); db.flush()
    db.add(PitReconciliationOrgSummary(workpaper_id=pre.id, company_id=company.id, period_id=period.id,
                                      org_code="17020", org_name="成都", org_full_name="成都分公司", declared_tax_amount=Decimal("100.00")))
    db.add(PitReconciliationOrgSummary(
        workpaper_id=workpaper.id, company_id=company.id, period_id=period.id,
        org_code="17020", org_name="成都", org_full_name="成都分公司",
        declared_tax_amount=Decimal("100.00"), certificate_tax_amount=Decimal("99.00"),
        difference_4=Decimal("1.00"), difference_4_manual_reason="凭证原因11",
        bank_tax_amount=Decimal("98.00"), difference_5=Decimal("1.00"), difference_5_manual_reason="银行原因12",
    ))
    db.add(PitBankTaxMatch(
        workpaper_id=workpaper.id, company_id=company.id, period_id=period.id,
        org_code="17020", org_full_name="成都分公司", bank_account="6222", transaction_time="2026-07-01",
        transaction_summary="个税扣款", debit_amount=Decimal("98.00"),
    ))
    source = tmp_path / "certificate.pdf"
    source.write_bytes(b"real-source-pdf")
    batch = ReconciliationImportBatch(
        company_id=company.id, period_id=period.id, import_type="tax_certificate",
        original_name="完税证明批量导入", stored_path=str(source), row_count=1,
        file_results=[{"file_name": "17020_综合所得申报.pdf", "status": "success", "fmss_file_id": 26}],
    )
    db.add(batch)
    db.flush()
    db.add(ReconciliationImportRow(
        company_id=company.id, period_id=period.id, batch_id=batch.id, import_type="tax_certificate", row_number=1,
        organization_code="17020", counterparty="成都分公司", amount=Decimal("99.00"), tax_period="2026-07",
        raw_data={"file_name": "17020_综合所得申报.pdf", "taxpayer_name": "成都分公司", "taxpayer_id": "9130", "tax_type": "个人所得税", "income_item": "工资薪金", "tax_period": "2026-07", "payment_date": "2026-07-01", "amount": "99.00"},
    ))
    db.commit()
    return db, company, period, workpaper


def test_post_builder_uses_ordered_columns_keeps_reason12_and_never_exports_raw_bank_rows(tmp_path):
    db, company, period, workpaper = _post_context(tmp_path)
    rows = PitFmssMapper().map_post(db, company.id, period.id, workpaper.id, _sheets())
    assert set(rows) == set(POST_SHEET_KEYS)
    assert rows["iit_post_tax_check"] == [["17020", "成都分公司", 100.0, 99.0, 1.0, "凭证原因11", 98.0, 1.0, "银行原因12"]]
    assert rows["iit_post_bank_flow"] == []
    assert rows["iit_post_bank_tax"] == [["17020", "成都分公司", "6222", "2026-07-01", "个税扣款", 98.0]]
    assert rows["iit_post_certificate"][0][-1] == 26
    path = PitFmssExcelBuilder().build(_sheets(), rows, prefix="iit-post-test")
    workbook = None
    try:
        workbook = openpyxl.load_workbook(path, read_only=True)
        assert workbook.sheetnames == list(POST_SHEET_KEYS)
        assert list(next(workbook["iit_post_tax_check"].values)) == list(POST_HEADERS["iit_post_tax_check"])
        assert list(next(workbook["iit_post_certificate"].values))[-1] == "fileId"
        assert list(workbook["iit_post_certificate"].values)[1][-1] == 26
        assert workbook["iit_post_bank_flow"].max_row == 1
    finally:
        if workbook is not None:
            workbook.close()
        path.unlink(missing_ok=True)


def test_post_reason12_mapping_uses_the_ninth_post_template_column(tmp_path):
    db, company, period, workpaper = _post_context(tmp_path)
    rows = PitFmssMapper().map_post(db, company.id, period.id, workpaper.id, _sheets())
    assert rows["iit_post_tax_check"][0][8] == "银行原因12"


def test_post_certificate_file_id_is_reused_for_multiple_tax_rows(tmp_path):
    db, company, period, workpaper = _post_context(tmp_path)
    db.add(ReconciliationImportRow(
        company_id=company.id, period_id=period.id, batch_id=db.query(ReconciliationImportBatch).filter_by(period_id=period.id, import_type="tax_certificate").one().id,
        import_type="tax_certificate", row_number=2, organization_code="17020", counterparty="成都分公司", amount=Decimal("1.00"), tax_period="2026-07",
        raw_data={"file_name": "17020_综合所得申报.pdf", "taxpayer_name": "成都分公司", "taxpayer_id": "9130", "tax_type": "个人所得税", "income_item": "利息税", "tax_period": "2026-07", "payment_date": "2026-07-01", "amount": "1.00"},
    ))
    db.commit()
    rows = PitFmssMapper().map_post(db, company.id, period.id, workpaper.id, _sheets())
    assert [row[-1] for row in rows["iit_post_certificate"]] == [26, 26]


def test_post_uses_latest_upload_file_id_for_all_rows_of_the_same_pdf(tmp_path):
    db, company, period, workpaper = _post_context(tmp_path)
    old_row = db.query(ReconciliationImportRow).filter_by(import_type="tax_certificate").one()
    batch = ReconciliationImportBatch(
        company_id=company.id, period_id=period.id, import_type="tax_certificate",
        original_name="再次上传", stored_path="", row_count=2,
        file_results=[{"file_name": "17020_综合所得申报.pdf", "status": "success", "fmss_file_id": 99}],
    )
    db.add(batch); db.flush()
    for index in (1, 2):
        db.add(ReconciliationImportRow(
            company_id=company.id, period_id=period.id, batch_id=batch.id, import_type="tax_certificate",
            row_number=index, organization_code="17020", amount=Decimal("49.50"),
            raw_data={**old_row.raw_data, "amount": "49.50"},
        ))
    db.commit()
    rows = PitFmssMapper().map_post(db, company.id, period.id, workpaper.id, _sheets())
    assert [row[9] for row in rows["iit_post_certificate"]] == [99, 99]
    assert PitOnlineService(db, SimpleNamespace())._post_certificate_file_ids(workpaper) == {
        "17020_综合所得申报.pdf": 99,
    }


def test_post_checks_actual_attachment_bindings_in_addition_to_file_id_column():
    payload = {"document": {"id": 3}, "sheets": {"iit_post_certificate": [
        ["税收完税证明", "one.pdf", "成都", "9130", "个人所得税", "工资薪金", "2026-07", "2026-07-01", "99", "26"],
        ["税收完税证明", "one.pdf", "成都", "9130", "个人所得税", "利息税", "2026-07", "2026-07-01", "1", "26"],
    ]}, "attachmentBindings": [
        {"fileId": 26, "fileName": "one.pdf", "sheetCode": "iit_post_certificate", "rowNo": 1},
        {"fileId": 26, "fileName": "one.pdf", "sheetCode": "iit_post_certificate", "rowNo": 2},
    ]}
    assert PitOnlineService._assert_post_certificate_file_ids(payload, {"one.pdf": 26})["matched_rows"] == 2
    for bad in ([], payload["attachmentBindings"][:1],
                [payload["attachmentBindings"][0], {**payload["attachmentBindings"][1], "fileId": 99}],
                [payload["attachmentBindings"][0], {**payload["attachmentBindings"][1], "rowNo": 1}]):
        with pytest.raises(FmssImportFailed, match="FMSS完税凭证文件绑定失败，已停止提交审核"):
            PitOnlineService._assert_post_certificate_file_ids({**payload, "attachmentBindings": bad}, {"one.pdf": 26})


def test_post_mapper_blocks_legacy_totals_that_differ_from_pre(tmp_path):
    db, company, period, workpaper = _post_context(tmp_path)
    db.query(PitReconciliationOrgSummary).filter_by(workpaper_id=workpaper.id).one().declared_tax_amount = Decimal("999")
    db.commit()
    with pytest.raises(FmssStateConflict, match="与缴款前版本不一致"):
        PitFmssMapper().map_post(db, company.id, period.id, workpaper.id, _sheets())


def test_post_mapper_requires_recalculation_after_pre_revision_changes(tmp_path):
    db, company, period, workpaper = _post_context(tmp_path)
    pre = db.query(PitReconciliationWorkpaper).filter_by(stage="pre_payment").one()
    workpaper.source_snapshot_json = {"pre_payment_snapshot": {"workpaper_id": pre.id, "draft_revision": pre.draft_revision}}
    pre.draft_revision += 1
    db.commit()
    with pytest.raises(FmssStateConflict, match="提交审核不会自动重算"):
        PitFmssMapper().map_post(db, company.id, period.id, workpaper.id, _sheets())


def test_post_submission_requires_fmss_pre_approval(tmp_path):
    db, company, period, workpaper = _post_context(tmp_path)

    class Client:
        def declaration(self, _branch, _month, stage):
            return {"document": {"status": "DRAFT" if stage == "PRE" else "DRAFT"}}

    service = PitOnlineService(db, Client())
    with pytest.raises(FmssStateConflict, match="缴款前底稿尚未复核通过"):
        service.prepare_post_upload(workpaper)


def test_submit_does_not_block_on_local_data_preparation(tmp_path, monkeypatch):
    path = tmp_path / "pre.xlsx"
    path.write_bytes(b"xlsx")

    class Client:
        submitted = False

        def import_iit_workbook(self, _branch, _month, _stage, _path):
            return {"documents": [{"branchCode": "19020", "id": 2}], "sheets": {}}

        def submit_iit_declaration(self, declaration_id, reviewer):
            assert (declaration_id, reviewer) == ("2", "LIUWT")
            self.submitted = True

    class Database:
        def commit(self):
            pass

    client = Client()
    service = PitOnlineService(Database(), client)  # type: ignore[arg-type]
    workpaper = SimpleNamespace(
        company_id=1,
        period_id=2,
        stage="pre_payment",
        calculation_status="idle",
        data_status="incomplete",
        draft_revision=1,
        platform_submission_id=None,
    )
    monkeypatch.setattr(service, "_assert_write_enabled", lambda _stage: None)
    monkeypatch.setattr(service, "prepare_pre_upload", lambda _workpaper: path)
    monkeypatch.setattr(service, "_context", lambda *_args: (SimpleNamespace(code="17020"), SimpleNamespace(year=2026, month=7)))
    monkeypatch.setattr(service, "_assert_reviewer", lambda *_args: None)
    monkeypatch.setattr(service, "sync_workpaper", lambda _workpaper: {"status": "REVIEWING"})
    monkeypatch.setattr(
        service,
        "query_declaration",
        lambda *_args: ({"document": {"id": 2, "status": "REVIEWING" if client.submitted else "DRAFT"}}, "REVIEWING" if client.submitted else "DRAFT"),
    )

    result = service.submit_workpaper(workpaper, "LIUWT")

    assert client.submitted is True
    assert result["status"] == "REVIEWING"


def test_post_submit_pipeline_imports_draft_binds_certificates_then_reviews(tmp_path, monkeypatch):
    db, company, period, workpaper = _post_context(tmp_path)
    monkeypatch.setattr(settings, "fmss_write_enabled", True)
    monkeypatch.setattr(settings, "fmss_post_write_enabled", True)
    monkeypatch.setattr(settings, "fmss_production_write_enabled", True)

    class Client:
        def __init__(self):
            self.imported = False
            self.submitted = False
            self.declaration_calls = []

        def declaration(self, _branch, _month, stage):
            self.declaration_calls.append(stage)
            if stage == "PRE":
                return {"document": {"id": 2, "status": "APPROVED"}}
            return {"document": {"id": 3, "status": "REVIEWING" if self.submitted else "DRAFT", "sheets": {"iit_post_certificate": [["税收完税证明", "17020_综合所得申报.pdf", "成都分公司", "9130", "个人所得税", "工资薪金", "2026-07", "2026-07-01", "99.00", 26]]}}}

        def sheets(self, stage):
            assert stage == "POST"
            return [{"key": sheet.key, "title": sheet.title, "headers": list(sheet.headers)} for sheet in _sheets()]

        def import_iit_workbook(self, branch, month, stage, path):
            assert (branch, month, stage) == ("19020", "2026-07", "POST")
            workbook = openpyxl.load_workbook(path, read_only=True)
            try:
                assert workbook["iit_post_bank_flow"].max_row == 1
            finally:
                workbook.close()
            self.imported = True
            return {"documents": [{"branchCode": "19020", "id": 3}], "sheets": {"iit_post_tax_check": {"imported": 1, "skipped": 0}, "iit_post_certificate": {"imported": 1, "skipped": 0}, "iit_post_bank_flow": {"imported": 0, "skipped": 0}, "iit_post_bank_tax": {"imported": 1, "skipped": 0}}, "totalImported": 3}

        def reviewers(self, _branch, _month, stage):
            assert stage == "POST"
            return [{"username": "LIUWT"}]

        def submit_iit_declaration(self, declaration_id, reviewer):
            assert (declaration_id, reviewer) == ("3", "LIUWT")
            self.submitted = True
            return 20

    client = Client()
    result = PitOnlineService(db, client).submit_workpaper(workpaper, "LIUWT")
    assert client.imported is True
    assert result["status"] == "REVIEWING"
    assert workpaper.workflow_status == "submitted"
    assert client.declaration_calls == ["POST", "PRE", "POST", "POST"]


def test_double_review_completion_requires_both_fmss_stages(tmp_path):
    db, company, period, post = _post_context(tmp_path)
    pre = db.query(PitReconciliationWorkpaper).filter_by(
        company_id=company.id, period_id=period.id, stage="pre_payment",
    ).one()
    pre.workflow_status = "submitted"
    db.commit()

    class Client:
        def declaration(self, _branch, _month, stage):
            return {"document": {"id": 2 if stage == "PRE" else 3, "status": "APPROVED"}}

    result = PitOnlineService(db, Client()).sync_workpaper(post)
    assert result["status"] == "APPROVED"
    assert result["double_review_completed"] is True


@pytest.mark.parametrize("state, locked", [("RETURNED", False), ("APPROVED", True)])
def test_post_remote_status_controls_local_edit_lock(tmp_path, state, locked):
    db, company, period, workpaper = _post_context(tmp_path)
    workpaper.platform_submission_id = "3"
    db.commit()

    class Client:
        def declaration(self, _branch, _month, _stage):
            return {"document": {"id": 3, "status": state}}

    service = PitOnlineService(db, Client())
    if locked:
        with pytest.raises(FmssStateConflict, match="不能修改"):
            service.enforce_editable(workpaper)
    else:
        service.enforce_editable(workpaper)
        assert workpaper.workflow_status == "returned"
