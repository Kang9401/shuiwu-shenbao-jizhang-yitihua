from openpyxl import load_workbook

from app.services.personnel_review import (
    build_personnel_change_review_table,
    write_personnel_change_review_table,
)


def test_transfer_hire_row_reuses_departed_staff_identity():
    report = {
        "personnel_changes": {
            "items": [
                {
                    "change_type": "入职",
                    "name": "张三",
                    "id_number": "",
                    "org_code_from": "",
                    "org_code_to": "10302",
                    "employee_id": "",
                    "phone": "",
                    "hire_date": "",
                    "leave_date": "",
                },
                {
                    "change_type": "离职",
                    "name": "张三",
                    "id_number": "110101199001010011",
                    "org_code_from": "10301",
                    "org_code_to": "",
                    "employee_id": "1001",
                    "phone": "13800000000",
                    "hire_date": "",
                    "leave_date": "2025-02-28",
                },
            ]
        }
    }

    table = build_personnel_change_review_table(report, 2025, 3)
    hire_row = table[table["人员状态"] == "正常"].iloc[0]

    assert hire_row["证件号码"] == "110101199001010011"
    assert hire_row["手机号码"] == "13800000000"
    assert hire_row["员工编号"] == "1001"
    assert hire_row["任职受雇从业日期"] == "2025/03/01"


def test_pure_new_hire_stays_blank_when_no_departed_match():
    report = {
        "personnel_changes": {
            "items": [
                {
                    "change_type": "入职",
                    "name": "李四",
                    "id_number": "",
                    "org_code_from": "",
                    "org_code_to": "10302",
                    "employee_id": "",
                    "phone": "",
                    "hire_date": "",
                    "leave_date": "",
                },
            ]
        }
    }

    table = build_personnel_change_review_table(report, 2025, 3)
    hire_row = table.iloc[0]

    assert hire_row["证件号码"] == ""
    assert hire_row["手机号码"] == ""
    assert hire_row["员工编号"] == ""


def test_departure_is_exported_as_current_month_non_normal_change():
    report = {
        "personnel_changes": {
            "items": [
                {
                    "change_type": "离职",
                    "name": "王五",
                    "id_number": "110101199001010011",
                    "org_code_from": "10301",
                    "org_code_to": "",
                    "employee_id": "1001",
                    "phone": "",
                    "hire_date": "",
                    "leave_date": "",
                },
            ]
        }
    }

    table = build_personnel_change_review_table(report, 2025, 4)

    assert len(table) == 1
    assert table.iloc[0]["人员状态"] == "非正常"


def test_export_marks_missing_required_cells_yellow(tmp_path):
    report = {
        "personnel_changes": {
            "items": [
                {
                    "change_type": "入职",
                    "name": "李四",
                    "id_number": "",
                    "org_code_to": "",
                    "employee_id": "",
                    "phone": "",
                    "hire_date": "",
                },
            ]
        }
    }
    table = build_personnel_change_review_table(report, 2025, 4)
    output = tmp_path / "人员信息变动表-雇员.xlsx"

    write_personnel_change_review_table(table, output)

    worksheet = load_workbook(output).active
    columns = {cell.value: cell.column for cell in worksheet[1]}
    for column in ["证件类型", "证件号码", "员工编号", "机构代码(人员信息表)"]:
        assert worksheet.cell(row=2, column=columns[column]).fill.fgColor.rgb == "00FFFF00"


def test_review_table_contains_all_monthly_changes_not_only_unresolved_personnel():
    report = {
        "personnel_changes": {
            "items": [
                {
                    "change_type": "入职",
                    "name": "待补充人员",
                    "can_confirm": False,
                    "missing_fields": ["证件号码"],
                    "org_code_to": "10301",
                },
                {
                    "change_type": "离职",
                    "name": "信息完整人员",
                    "can_confirm": True,
                    "missing_fields": [],
                    "id_number": "110101199001010011",
                    "org_code_from": "10301",
                },
            ]
        }
    }

    table = build_personnel_change_review_table(report, 2025, 4)

    assert table["*姓名"].tolist() == ["待补充人员", "信息完整人员"]
