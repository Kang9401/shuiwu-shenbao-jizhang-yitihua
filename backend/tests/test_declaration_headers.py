from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.core.config import settings
from app.services.generation import generate_declarations
from app.services.verification import DECLARATION_COLUMNS
from app.workflows import get_workflow


def _source():
    row = {column: 0 for column in DECLARATION_COLUMNS}
    row.update({"工号": "001", "*姓名": "测试", "*证件类型": "居民身份证",
                "*证件号码": "001234567890123456", "住房公积金调整": 12.34,
                "机构代码": "17020"})
    return pd.DataFrame([row])


def _check(path, source):
    with Path(path).open("rb") as stream:
        book = load_workbook(stream)
    header = [cell.value for cell in book.active[1]]
    position = DECLARATION_COLUMNS.index("住房公积金调整")
    assert header[position] == "住房公积金和年金调整"
    assert "住房公积金调整" not in header
    assert header[position - 1] == "律师办案费用"
    assert header[position + 1] == "西藏附加减除费用"
    assert book.active.cell(2, position + 1).value == "12.34"
    assert source.loc[0, "住房公积金调整"] == 12.34
    assert "住房公积金和年金调整" not in source.columns


def test_general_workflow_exports_current_adjustment_header(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "storage_root", tmp_path)
    source = _source()
    artifacts = get_workflow("general_salary_tax")._write_split_artifacts(1, source, 2026, 9)
    _check(artifacts[0][1], source)


def test_salary_generation_exports_current_header_without_changing_snapshot(tmp_path, monkeypatch):
    # Exercise header serialization independently of the legacy XLS writer.
    original = pd.DataFrame.to_excel

    def write_xlsx(frame, path, **kwargs):
        with Path(path).open("wb") as stream:
            return original(frame, stream, engine="openpyxl", **kwargs)

    monkeypatch.setattr(pd.DataFrame, "to_excel", write_xlsx)
    source = _source()
    paths = generate_declarations(source, str(tmp_path), 2026, 9)
    _check(paths[0], source)
