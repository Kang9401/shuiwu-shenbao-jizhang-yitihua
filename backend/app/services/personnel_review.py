"""人员变动 HR 复核底稿。

把核对报告中的入职、离职、调岗结果整理成旧流程里的
`人员信息变动表-雇员.xlsx`，供 HR 补充/确认后再导入人员信息更新流程。
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import PatternFill


PERSONNEL_CHANGE_COLUMNS = [
    "*姓名", "证件类型", "证件号码", "国籍(地区)", "性别", "出生日期",
    "人员状态", "任职受雇从业类型", "手机号码", "任职受雇从业日期",
    "离职日期", "机构代码(人员信息表)", "机构代码(工资单)", "员工编号",
    "涉税事由", "出生国家(地区)",
]

MISSING_REQUIRED_FILL = PatternFill(fill_type="solid", fgColor="FFFF00")


def _clean(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "nat"} else text


def _id_gender(id_number: str) -> str:
    text = _clean(id_number)
    if len(text) != 18 or not text[-2].isdigit():
        return ""
    return "男" if int(text[-2]) % 2 == 1 else "女"


def _id_birth_date(id_number: str) -> str:
    text = _clean(id_number)
    if len(text) != 18 or not text[6:14].isdigit():
        return ""
    return f"{text[6:10]}-{text[10:12]}-{text[12:14]}"


def build_personnel_change_review_table(report: dict, year: int, month: int) -> pd.DataFrame:
    """生成包含本月正常与非正常变动的雇员信息表。"""
    rows: list[dict[str, Any]] = []
    default_hire_date = f"{year}/{month:02d}/01"
    all_change_items = report.get("personnel_changes", {}).get("items", [])
    # Export every detected monthly change. This is a change table, never a full employee roster.
    # Missing fields remain highlighted so HR can identify the rows that require completion.
    change_items = all_change_items

    departed_by_name: dict[str, dict[str, Any] | None] = {}
    for item in change_items:
        if item.get("change_type") != "离职":
            continue
        name = _clean(item.get("name"))
        if not name:
            continue
        departed_by_name[name] = None if name in departed_by_name else item

    def append_row(item: dict, status: str, org_staff: str, org_payroll: str, hire_date: str, leave_date: str) -> None:
        id_number = _clean(item.get("id_number"))
        rows.append({
            "*姓名": _clean(item.get("name")),
            "证件类型": "居民身份证" if id_number else "",
            "证件号码": id_number,
            "国籍(地区)": "中国" if id_number else "",
            "性别": _id_gender(id_number),
            "出生日期": _id_birth_date(id_number),
            "人员状态": status,
            "任职受雇从业类型": "雇员",
            "手机号码": _clean(item.get("phone")),
            "任职受雇从业日期": hire_date,
            "离职日期": leave_date,
            "机构代码(人员信息表)": org_staff,
            "机构代码(工资单)": org_payroll,
            "员工编号": _clean(item.get("employee_id")),
            "涉税事由": _clean(item.get("tax_reason")) or ("其他" if _clean(item.get("cert_type")) not in {"", "居民身份证"} else ""),
            "出生国家(地区)": _clean(item.get("birth_country")) or ("国籍" if _clean(item.get("cert_type")) not in {"", "居民身份证"} else ""),
        })

    def enrich_transfer_hire(item: dict) -> dict:
        """同名唯一离职 + 入职视为疑似变动单位，入职行沿用旧人员基础信息。"""
        departed = departed_by_name.get(_clean(item.get("name")))
        if not departed:
            return item
        enriched = dict(item)
        for field in ["id_number", "phone"]:
            if not _clean(enriched.get(field)):
                enriched[field] = departed.get(field, "")
        if not _clean(enriched.get("employee_id")):
            enriched["employee_id"] = departed.get("employee_id", "")
        return enriched

    for item in change_items:
        change_type = item.get("change_type")
        if change_type == "入职":
            item = enrich_transfer_hire(item)
            org = _clean(item.get("org_code_to"))
            append_row(item, "正常", org, org, _clean(item.get("hire_date")) or default_hire_date, "")
        elif change_type == "离职":
            append_row(
                item,
                "非正常",
                _clean(item.get("org_code_from")),
                "",
                _clean(item.get("hire_date")),
                _clean(item.get("leave_date")),
            )
        elif change_type == "调岗":
            append_row(
                item,
                "非正常",
                _clean(item.get("org_code_from")),
                "",
                _clean(item.get("hire_date")),
                _clean(item.get("leave_date")),
            )
            new_org = _clean(item.get("org_code_to"))
            append_row(item, "正常", new_org, new_org, default_hire_date, "")

    return pd.DataFrame(rows, columns=PERSONNEL_CHANGE_COLUMNS)


def write_personnel_change_review_table(df: pd.DataFrame, path: str | Path) -> None:
    """导出雇员变动表，并用黄色标记缺失的必填字段。"""
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="人员信息变动表-雇员")
        worksheet = writer.sheets["人员信息变动表-雇员"]
        column_index = {cell.value: cell.column for cell in worksheet[1]}

        for row_number, (_, row) in enumerate(df.iterrows(), start=2):
            required = ["*姓名", "证件类型", "证件号码", "员工编号", "机构代码(人员信息表)"]
            if _clean(row.get("人员状态")) == "非正常":
                required.extend(["手机号码", "任职受雇从业日期"])
            cert_type = _clean(row.get("证件类型"))
            if cert_type and cert_type != "居民身份证":
                required.extend(["国籍(地区)", "性别", "出生日期", "涉税事由", "出生国家(地区)"])
            for column in required:
                if not _clean(row.get(column)):
                    worksheet.cell(row=row_number, column=column_index[column]).fill = MISSING_REQUIRED_FILL
