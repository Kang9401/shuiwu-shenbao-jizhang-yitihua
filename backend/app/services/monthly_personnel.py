"""Shared review lifecycle for broker, intern and labor remuneration workflows."""
from __future__ import annotations

import calendar
from datetime import datetime
from pathlib import Path

import pandas as pd

from app.models.accounting import OrganizationMapping
from app.models.core import Artifact, Job, Period, UploadedFile
from app.services.excel import write_workbook
from app.services.excel import read_excel
from app.services.personnel_fields import text, with_foreign_defaults, missing_personnel_fields
from app.services.personnel_master import PersonnelMasterResolver
from app.services.personnel_review import PERSONNEL_CHANGE_COLUMNS
from app.services.personnel_update import build_personnel_collection_files
from app.services.storage import artifact_path
from app.workflows.base import WorkflowResult

SPECS = {
    "broker_tax": ("broker", "经纪人", "broker_income", "broker_tax_result"),
    "intern_tax": ("intern", "实习生", "intern_salary", "intern_tax_result"),
    "part_time_tax": ("part_time", "劳务报酬", "payroll", "part_time_workpaper"),
}
ALIASES = {
    "*姓名": ["*姓名", "姓名", "员工姓名", "经纪人姓名"],
    "员工编号": ["员工编号", "工号", "人员编号", "员工编码"],
    "证件类型": ["证件类型", "*证件类型"],
    "证件号码": ["证件号码", "*证件号码", "证件号码（同名必填）", "身份证号码"],
    "国籍(地区)": ["国籍(地区)", "国籍（地区）", "*国籍(地区)", "*国籍（地区）", "国籍"],
    "性别": ["性别", "*性别"], "出生日期": ["出生日期", "*出生日期"],
    "手机号码": ["手机号码", "*手机号码", "联系方式"],
    "任职受雇从业日期": ["任职受雇从业日期", "任职/从业日期", "*任职/从业日期", "实习开始时间"],
    "机构代码(人员信息表)": ["机构代码(人员信息表)", "机构代码", "分支机构代码", "归属机构代码", "单位编号"],
    "人员状态": ["人员状态", "状态"], "离职日期": ["离职日期"],
    "涉税事由": ["涉税事由"], "出生国家(地区)": ["出生国家(地区)", "出生国家（地区）"],
}


def normalize(frame):
    result = []
    for raw in frame.fillna("").to_dict("records"):
        record = {key: next((text(raw.get(c)) for c in names if text(raw.get(c))), "") for key, names in ALIASES.items()}
        if not any(record.get(key) for key in ("*姓名", "员工编号", "证件号码")):
            continue
        for key in ("员工编号", "机构代码(人员信息表)"):
            if record[key].endswith(".0"):
                record[key] = record[key][:-2]
        record["人员状态"] = record["人员状态"] or "正常"
        if record["证件类型"] == "身份证":
            record["证件类型"] = "居民身份证"
        result.append(with_foreign_defaults({**raw, **record}))
    return result


def reconcile_people(master, payroll, year, month, changes=None):
    """Absence of a row means departure; zero income is still a present person."""
    master = [dict(row) for row in master]
    issues, actions, matches, seen = [], [], [], set()
    for position, person in enumerate(payroll):
        candidates = []
        for field in ("员工编号", "证件号码", "*姓名"):
            value = text(person.get(field))
            if value:
                candidates = [i for i, row in enumerate(master) if text(row.get(field)) == value]
                if candidates:
                    break
        if len(candidates) > 1:
            issues.append({"message": f"{person['*姓名']} 无法唯一匹配人员主数据", "issue_type": "ambiguous_person"})
            matches.append(None)
            continue
        if candidates:
            idx = candidates[0]
            if idx in seen:
                issues.append({"message": f"{person['*姓名']} 本月收入记录重复，请合并后上传", "issue_type": "duplicate_income"})
            seen.add(idx)
            row = master[idx]
            action = "重新任职" if row["人员状态"] not in {"正常", "在职"} else ""
            if action:
                row["人员状态"], row["离职日期"] = "正常", ""
                row["任职受雇从业日期"] = text(person.get("任职受雇从业日期")) or f"{year}-{month:02d}-01"
            if action or missing_personnel_fields(row):
                actions.append((idx, action or "资料补充"))
        else:
            idx = len(master)
            row = with_foreign_defaults(dict(person))
            row["人员状态"] = "正常"
            row["任职受雇从业日期"] = text(row.get("任职受雇从业日期")) or f"{year}-{month:02d}-01"
            master.append(row)
            seen.add(idx)
            actions.append((idx, "新增"))
        matches.append(idx)
    # Do not infer departures from an invalid/ambiguous input batch.
    if not issues and payroll:
        for idx, row in enumerate(master):
            if idx not in seen and row["人员状态"] in {"正常", "在职"}:
                row["人员状态"] = "非正常"
                from datetime import date, timedelta
                row["离职日期"] = (date(year, month, 1) - timedelta(days=1)).isoformat()
                actions.append((idx, "离职"))
    action_ids = {str(idx + 1): idx for idx, _ in actions}
    updated = set()
    for change in (changes or []):
        key = text(change.get("核对行号"))
        if key.endswith(".0"):
            key = key[:-2]
        if key not in action_ids or key in updated:
            issues.append({"issue_type": "invalid_change", "message": "变更表核对行号无效或重复，请使用本轮下载的变更表"})
            continue
        updated.add(key)
        row = master[action_ids[key]]
        for column in PERSONNEL_CHANGE_COLUMNS:
            if column not in {"人员状态", "机构代码(工资单)"} and text(change.get(column)):
                row[column] = text(change[column])
        row.update(with_foreign_defaults(row))
    review = []
    for idx, action in actions:
        row = master[idx]
        missing = missing_personnel_fields(row)
        if missing:
            issues.append({"issue_type": "missing_personnel_fields", "message": f"{row['*姓名']} 缺少：{'、'.join(missing)}", "missing_fields": missing})
        review.append({"核对行号": str(idx+1), "变更类型": action, **{c: row.get(c, "") for c in PERSONNEL_CHANGE_COLUMNS}})
    return master, review, matches, issues


def _write_review(path, records):
    from openpyxl.styles import PatternFill
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.utils import get_column_letter
    columns = ["核对行号", "变更类型", *PERSONNEL_CHANGE_COLUMNS]
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(records, columns=columns).to_excel(writer, index=False, sheet_name="人员信息变更表")
        ws = writer.sheets["人员信息变更表"]
        for r, record in enumerate(records, 2):
            for name in missing_personnel_fields(record):
                cell = f"{get_column_letter(columns.index(name)+1)}{r}"
                ws.conditional_formatting.add(cell, FormulaRule(formula=[f'LEN(TRIM({cell}))=0'], fill=PatternFill('solid',fgColor='FFFF00')))
                if not str(record.get(name, '') or '').strip():
                    ws[cell].fill = PatternFill('solid', fgColor='FFFF00')


def run_monthly_personnel(workflow, db, job_id, period_id, files, operation):
    person_type, label, role, result_type = SPECS[workflow.info.code]
    period = db.query(Period).filter_by(id=period_id).one()
    # Use a frozen baseline throughout this round. Generating never reuses a newly saved master as its baseline.
    initial = db.query(Job).filter_by(period_id=period_id, workflow_code=workflow.info.code, operation="initial").filter(Job.id != job_id).order_by(Job.id.desc()).first()
    source_files = [f for f in files if f.file_role == role]
    # Older browser sessions may have uploaded the file with the workflow's
    # legacy role name.  The monthly personnel screen has a single income
    # upload, so safely treat those files as the income source when no exact
    # role match is present.
    if not source_files and files and operation == "initial":
        source_files = list(files)
    if operation != "initial" and not source_files and initial:
        source_files = db.query(UploadedFile).filter(UploadedFile.id.in_(initial.input_file_ids), UploadedFile.file_role == role).all()
        if not source_files:
            source_files = db.query(UploadedFile).filter(UploadedFile.id.in_(initial.input_file_ids)).all()
    root = artifact_path(job_id, "review").parent
    root.mkdir(parents=True, exist_ok=True)
    snapshot = None
    if initial and operation != "initial":
        snapshot = db.query(Artifact).filter_by(job_id=initial.id, artifact_type="personnel_baseline").first()
    if snapshot:
        source_master = pd.read_excel(snapshot.stored_path, dtype=str).fillna("")
    else:
        # Each monthly comparison starts from the previous month's frozen
        # personnel master. The current month's income sheet supplies only
        # this month's presence and income changes.
        y, m = (period.year, period.month - 1) if period.month > 1 else (period.year - 1, 12)
        previous = db.query(Period).filter_by(company_id=period.company_id, year=y, month=m).first()
        path = PersonnelMasterResolver(db, previous.id).resolve_path(person_type) if previous else None
        # Bootstrap the very first configured period when no prior period
        # exists; all subsequent periods are strictly previous-month based.
        if path is None and previous is None:
            path = PersonnelMasterResolver(db, period_id).resolve_path(person_type)
        source_master = pd.read_excel(path,dtype=str).fillna("") if path else pd.DataFrame(columns=PERSONNEL_CHANGE_COLUMNS)
    baseline = root / f"{label}_本轮人员基准.xlsx"
    write_workbook(baseline, {"人员信息表":source_master},text_values=True)
    if person_type == "intern":
        frames = workflow._prepare_frames(db, workflow._frames_with_aliases(source_files))
        raw = frames.get("intern_salary_sheet",pd.DataFrame())
    elif person_type == "broker":
        # This screen has exactly one income input. Read it directly so an
        # old role alias or an xls workbook sheet name cannot hide the data.
        raw = pd.concat([read_excel(f.stored_path) for f in source_files], ignore_index=True, sort=False) if source_files else pd.DataFrame()
        frames = {"broker_salary_sheet": raw}
    else:
        raw = pd.concat([pd.read_excel(f.stored_path,dtype=str) for f in source_files],ignore_index=True) if source_files else pd.DataFrame()
        frames = {}
    payroll = normalize(raw)
    org_names = {r.branch_name:r.org_code for r in db.query(OrganizationMapping).filter_by(active=1)}
    for row in payroll:
        if not row["机构代码(人员信息表)"]:
            row["机构代码(人员信息表)"] = org_names.get(text(row.get("营业部名称")), "")
    change_files = [f for f in files if f.file_role == "personnel_changes"]
    if operation in {"recheck", "generate"} and not change_files and initial:
        latest = db.query(Job).filter_by(period_id=period_id,workflow_code=workflow.info.code,operation="recheck").filter(Job.id>initial.id, Job.id != job_id).order_by(Job.id.desc()).first()
        if latest:
            change_files = db.query(UploadedFile).filter(UploadedFile.id.in_(latest.input_file_ids),UploadedFile.file_role=="personnel_changes").all()
    changes = [row for f in change_files for row in pd.read_excel(f.stored_path,dtype=str).fillna("").to_dict("records")]
    master, review, matches, issues = reconcile_people(normalize(source_master), payroll, period.year,period.month,changes)
    if operation not in {"initial", "recheck", "generate"}:
        issues.append({"issue_type":"invalid_operation", "message":"不支持的核对操作"})
    if operation != "initial" and (initial is None or snapshot is None):
        issues.append({"issue_type":"missing_initial", "message":"请先上传本月收入资料并执行首轮核对"})
    if not payroll:
        issues.append({"issue_type":"missing_income","message":"请上传包含人员记录的本月收入资料"})
    for row in master:
        if row.get("机构代码(人员信息表)") and row["机构代码(人员信息表)"] not in org_names.values():
            issues.append({"issue_type":"unknown_org","message":f"{row['*姓名']} 的机构代码未维护"})
    detail = pd.DataFrame()
    if not issues:
        if person_type == "broker":
            frames["broker_staff_info"] = pd.DataFrame([{**r, "机构代码":r["机构代码(人员信息表)"]} for r in master]).drop(columns=["机构代码(人员信息表)"], errors="ignore")
            result = workflow.transform(frames)
        elif person_type == "intern":
            enriched = []
            for person, idx in zip(payroll,matches):
                row = {**person, **{c:master[idx].get(c, "") for c in ALIASES}}
                for target,source in {"*证件类型":"证件类型","*证件号码":"证件号码","*国籍(地区)":"国籍(地区)","*性别":"性别","*出生日期":"出生日期","机构代码":"机构代码(人员信息表)","实习开始时间":"任职受雇从业日期"}.items():
                    row[target] = row.get(source,"")
                enriched.append(row)
            frames["intern_salary_sheet"] = pd.DataFrame(enriched)
            result = workflow.transform(frames)
        else:
            from app.services.part_time_tax import _normalize_payroll, _number, DECLARATION_COLUMNS, PART_TIME_INCOME_ITEM
            income = _normalize_payroll(pd.DataFrame(payroll))
            rows = []
            for (_, source), idx in zip(income.iterrows(),matches):
                row = master[idx]
                amount = pd.to_numeric(source["本期收入"], errors="coerce")
                # A zero-income row still proves that the person is present this month.
                # Keep the row for downstream salary-style processing; only a non-numeric
                # amount is invalid input.
                if pd.isna(amount):
                    issues.append({"issue_type":"invalid_income","message":f"{row['*姓名']} 本期收入必须为数字"})
                rows.append({**{c:_number(source.get(c)) for c in DECLARATION_COLUMNS},"工号":row.get("员工编号",""),"*姓名":row['*姓名'],"*证件类型":row['证件类型'],"*证件号码":row['证件号码'],"*所得项目":PART_TIME_INCOME_ITEM,"备注":source.get("备注",""),"归属机构代码":row['机构代码(人员信息表)']})
            result = {"detail":pd.DataFrame(rows),"issues":[]}
        detail = result.get("detail",pd.DataFrame())
        issues.extend(result.get("issues",[]))
    blocking = any(i.get("severity") != "warning" for i in issues)
    review_path = root / f"{label}_人员信息变更表.xlsx"
    _write_review(review_path,review)
    output = root / f"{label}_核对底稿.xlsx"
    write_workbook(output,{"申报明细":detail,"人员变化":pd.DataFrame(review),"问题清单":pd.DataFrame(issues)})
    artifacts = [("personnel_baseline",str(baseline)),("personnel_change_review",str(review_path)),(result_type,str(output))]
    finalized = operation == "generate" and not blocking
    if finalized:
        collection = pd.DataFrame(review).drop(columns=["核对行号","变更类型"],errors="ignore")
        if not collection.empty:
            collection["任职受雇从业类型"] = {"intern":"实习学生（全日制学历教育）","broker":"其他","part_time":"其他"}[person_type]
            artifacts += [("personnel_collection",p) for p in build_personnel_collection_files(collection,root,period.year,period.month,person_type_label=label,duplicate_keep="first",file_extension=".xlsx")]
        if person_type == "part_time":
            from app.services.part_time_tax import DECLARATION_COLUMNS
            for org, group in detail.groupby("归属机构代码"):
                path = root / f"{org}_劳务报酬所得_{period.year}{period.month:02d}.xlsx"
                write_workbook(path,{"Sheet1":group.reindex(columns=DECLARATION_COLUMNS)},text_values=True)
                artifacts.append(("declaration",str(path)))
        else:
            artifacts += workflow._write_split_artifacts(job_id,detail,period.year,period.month)
    return WorkflowResult(status="needs_review" if blocking else "success",summary={"operation":operation,"rows":len(detail),"issues":len(issues),"issue_details":issues,"can_generate":not blocking,"finalized":finalized,"new_count":sum(r['变更类型']=='新增' for r in review),"leaver_count":sum(r['变更类型']=='离职' for r in review),"matched_count":sum(i is not None for i in matches),"payroll_count":len(payroll)},artifact_paths=artifacts)
