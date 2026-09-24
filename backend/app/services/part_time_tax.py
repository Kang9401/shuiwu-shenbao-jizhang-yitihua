from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
import re
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from app.core.config import settings
from app.models.accounting import OrganizationMapping, PersonnelMasterArtifact, PersonnelMasterImportBatch
from app.models.core import Period
from app.services.excel import read_excel, write_workbook
from app.services.personnel_master import PersonnelMasterResolver
from app.services.storage import artifact_path


PART_TIME_INCOME_ITEM = "一般劳务报酬所得"
TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "resources" / "非全日制用工_三张模板示例.xlsx"
TEMPLATE_DIR = TEMPLATE_PATH.parent
SOURCE_TEMPLATE_PATHS = {
    "人员信息表": TEMPLATE_DIR / "非全日制人员信息表模板.xlsx",
    "人员信息变更表": TEMPLATE_DIR / "非全日制人员信息变更表模板.xlsx",
}
MASTER_COLUMNS = [
    "姓名", "工号", "归属机构代码", "归属机构名称", "证件类型", "证件号码",
    "国籍（地区）", "性别", "出生日期", "手机号码", "任职/从业日期", "人员状态", "离职日期",
]
CHANGE_COLUMNS = [
    "变更类型", "姓名", "工号", "归属机构代码", "归属机构名称", "证件类型", "证件号码",
    "国籍（地区）", "性别", "出生日期", "手机号码", "任职/从业日期", "离职日期", "备注",
]
PAYROLL_COLUMNS = [
    "姓名", "证件号码（同名必填）", "本期收入", "本期免税收入", "商业健康保险",
    "税延养老保险", "其他", "允许扣除的税费", "减免税额", "协定减免", "备注",
]
DECLARATION_COLUMNS = [
    "工号", "*姓名", "*证件类型", "*证件号码", "*所得项目", "本期收入", "本期免税收入",
    "商业健康保险", "税延养老保险", "其他", "允许扣除的税费", "减免税额", "协定减免", "备注",
]


def _clean(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "nat"} else text


def _first(row: pd.Series, names: list[str]) -> str:
    for name in names:
        if name in row.index and _clean(row.get(name)):
            return _clean(row.get(name))
    return ""


def _normalize_header(value: Any) -> str:
    """Normalize template-required markers before matching source columns."""
    text = _clean(value).replace("＊", "*")
    return text.lstrip("*").strip()


def _normalize_headers(frame: pd.DataFrame) -> pd.DataFrame:
    """Return a canonical-header view while preserving the first non-empty duplicate."""
    result = pd.DataFrame(index=frame.index)
    for column in frame.columns:
        canonical = _normalize_header(column)
        if not canonical:
            continue
        if canonical not in result.columns:
            result[canonical] = frame[column]
        else:
            existing = result[canonical].map(_clean)
            result[canonical] = result[canonical].where(existing != "", frame[column])
    return result


def normalize_part_time_name(value: Any) -> str:
    return _clean(value).replace(" ", "")


def _number(value: Any, default: float = 0.0) -> float:
    text = _clean(value).replace(",", "")
    if not text:
        return default
    try:
        return float(text)
    except ValueError:
        return default


def _date_text(value: Any) -> str:
    if not _clean(value):
        return ""
    parsed = pd.to_datetime(value, errors="coerce")
    return "" if pd.isna(parsed) else parsed.strftime("%Y-%m-%d")


def _valid_id(value: str) -> bool:
    value = value.upper()
    if len(value) != 18 or not value[:17].isdigit() or not (value[17].isdigit() or value[17] == "X"):
        return False
    weights = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
    checks = "10X98765432"
    return checks[sum(int(a) * b for a, b in zip(value[:17], weights)) % 11] == value[17]


def _previous_period(db, period: Period | None) -> Period | None:
    if period is None:
        return None
    year = period.year if period.month > 1 else period.year - 1
    month = period.month - 1 if period.month > 1 else 12
    return db.query(Period).filter(Period.year == year, Period.month == month).first()


def _empty_master() -> pd.DataFrame:
    return pd.DataFrame(columns=MASTER_COLUMNS)


def load_previous_master(db, period_id: int) -> pd.DataFrame:
    return _load_previous_master_source(db, period_id)[0]


def _load_previous_master_source(db, period_id: int) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Load only the immediately preceding month's part-time master."""
    period = db.query(Period).filter(Period.id == period_id).first()
    previous = _previous_period(db, period)
    metadata: dict[str, Any] = {
        "period": f"{previous.year}{previous.month:02d}" if previous else "",
        "file": "",
        "path": "",
        "missing": previous is None,
        "month_warning": "",
    }
    if previous is None:
        return _empty_master(), metadata
    path = PersonnelMasterResolver(db, previous.id).resolve_path("part_time")
    if not path or not Path(path).exists():
        metadata["missing"] = True
        return _empty_master(), metadata
    metadata["missing"] = False
    metadata["path"] = str(path)
    metadata["file"] = Path(path).name
    filename_month = re.search(r"(20\d{2})(0[1-9]|1[0-2])", Path(path).name)
    expected_month = metadata["period"]
    if filename_month and expected_month and "".join(filename_month.groups()) != expected_month:
        metadata["month_warning"] = (
            f"主数据文件名月份为 {''.join(filename_month.groups())}，当前按上月 {expected_month} 主数据处理"
        )
    source = read_excel(path).fillna("")
    source = _normalize_headers(source)
    result = pd.DataFrame(index=source.index)
    for column in MASTER_COLUMNS:
        aliases = {
            "姓名": ["姓名"], "工号": ["工号", "员工编号"],
            "归属机构代码": ["归属机构代码", "机构代码", "机构代码(人员信息表)"],
            "归属机构名称": ["归属机构名称", "机构名称"], "证件类型": ["证件类型"],
            "证件号码": ["证件号码"], "国籍（地区）": ["国籍（地区）", "国籍(地区)"],
            "性别": ["性别"], "出生日期": ["出生日期"],
            "手机号码": ["手机号码"], "任职/从业日期": ["任职/从业日期", "任职受雇从业日期"],
            "人员状态": ["人员状态", "状态"], "离职日期": ["离职日期"],
        }[column]
        found = next((name for name in aliases if name in source.columns), None)
        result[column] = source[found].map(_clean) if found else ""
    result["人员状态"] = result["人员状态"].replace("", "正常")
    return result, metadata


def materialize_previous_part_time_master(db, period_id: int) -> dict[str, Any]:
    """Copy the previous month's canonical source into this month's monthly artifact."""
    master, metadata = _load_previous_master_source(db, period_id)
    if metadata["missing"]:
        return metadata
    period = db.query(Period).filter(Period.id == period_id).first()
    if period is None:
        return metadata
    target_dir = settings.artifact_dir / "personnel_masters" / str(period_id) / "part_time"
    target_dir.mkdir(parents=True, exist_ok=True)
    target_path = target_dir / f"{period.year}{period.month:02d}_part_time_month_all.xlsx"
    write_workbook(target_path, {"人员信息表": master}, text_values=True)
    artifact = (
        db.query(PersonnelMasterArtifact)
        .filter(
            PersonnelMasterArtifact.period_id == period_id,
            PersonnelMasterArtifact.person_type == "part_time",
            PersonnelMasterArtifact.scope_type == "month",
            PersonnelMasterArtifact.scope_code == "",
        )
        .first()
    )
    if artifact is None:
        artifact = PersonnelMasterArtifact(
            period_id=period_id, person_type="part_time", scope_type="month", scope_code="",
            file_name=target_path.name, stored_path=str(target_path), row_count=len(master), validation_issues=[],
        )
        db.add(artifact)
    else:
        artifact.file_name = target_path.name
        artifact.stored_path = str(target_path)
        artifact.row_count = len(master)
        artifact.validation_issues = []
        artifact.updated_at = datetime.utcnow()
    db.flush()
    db.add(PersonnelMasterImportBatch(
        period_id=period_id,
        artifact_id=artifact.id,
        person_type="part_time",
        scope_type="month",
        scope_code="",
        original_name=f"自动复制{metadata['period']}非全日制人员主数据：{metadata['file']}",
        stored_path=str(target_path),
        row_count=len(master),
        validation_issues=[],
    ))
    db.commit()
    metadata["materialized"] = True
    metadata["target_path"] = str(target_path)
    return metadata


def _normalize_changes(frame: pd.DataFrame) -> pd.DataFrame:
    frame = _normalize_headers(frame)
    result = pd.DataFrame(index=frame.index)
    aliases = {
        "变更类型": ["变更类型", "变更/调整类型"], "姓名": ["姓名"], "工号": ["工号", "员工编号"],
        "归属机构代码": ["归属机构代码", "机构代码"], "归属机构名称": ["归属机构名称", "机构名称"],
        "证件类型": ["证件类型", "*证件类型"], "证件号码": ["证件号码", "*证件号码"],
        "国籍（地区）": ["国籍（地区）", "国籍(地区)", "*国籍(地区)"], "性别": ["性别", "*性别"],
        "出生日期": ["出生日期", "*出生日期"], "手机号码": ["手机号码"],
        "任职/从业日期": ["任职/从业日期", "任职受雇从业日期"], "离职日期": ["离职日期"], "备注": ["备注"],
    }
    for column, candidates in aliases.items():
        found = next((name for name in candidates if name in frame.columns), None)
        result[column] = frame[found].map(_clean) if found else ""
    return result


def _normalize_payroll(frame: pd.DataFrame) -> pd.DataFrame:
    frame = _normalize_headers(frame)
    aliases = {
        "姓名": ["姓名"], "证件号码": ["证件号码（同名必填）", "证件号码"],
        "本期收入": ["本期收入", "收入"], "本期免税收入": ["本期免税收入", "免税收入"],
        "商业健康保险": ["商业健康保险"], "税延养老保险": ["税延养老保险"], "其他": ["其他"],
        "允许扣除的税费": ["允许扣除的税费"], "减免税额": ["减免税额"], "协定减免": ["协定减免"], "备注": ["备注"],
    }
    result = pd.DataFrame(index=frame.index)
    for column, candidates in aliases.items():
        found = next((name for name in candidates if name in frame.columns), None)
        result[column] = frame[found].map(_clean) if found else ""
    return result


def _org_names(db) -> dict[str, str]:
    return {item.org_code.strip(): item.branch_name.strip() for item in db.query(OrganizationMapping).filter(OrganizationMapping.active == 1).all()}


def _apply_change(master: pd.DataFrame, row: pd.Series, org_names: dict[str, str], issues: list[dict[str, Any]]) -> None:
    name = normalize_part_time_name(row["姓名"])
    change_type = _clean(row["变更类型"])
    id_number = _clean(row["证件号码"])
    matches = master.index[master["姓名"].map(normalize_part_time_name) == name]
    if id_number and len(matches) > 1:
        matches = master.index[master["证件号码"].map(_clean) == id_number]
    if change_type in {"新增", "重新任职"}:
        required = ["归属机构代码", "证件类型", "证件号码", "国籍（地区）", "性别", "出生日期", "手机号码", "任职/从业日期"]
        missing = [column for column in required if not _clean(row[column])]
        if missing:
            issues.append({"issue_type": "part_time_change_missing", "message": f"{name} 的变更资料缺少：{'、'.join(missing)}", "severity": "error"})
            return
        if _clean(row["证件类型"]) == "居民身份证" and not _valid_id(id_number):
            issues.append({"issue_type": "invalid_id_number", "message": f"{name} 的居民身份证号码校验不通过"})
            return
        if change_type == "新增" and len(matches):
            issues.append({"issue_type": "duplicate_part_time_addition", "message": f"{name} 已存在于非全日制人员主数据"})
            return
        if change_type == "重新任职" and len(matches) != 1:
            issues.append({"issue_type": "rehire_not_matched", "message": f"{name} 重新任职无法唯一匹配历史人员"})
            return
        target = matches[0] if len(matches) == 1 else len(master)
        values = {column: _clean(row[column]) for column in MASTER_COLUMNS if column in row.index}
        values["归属机构名称"] = values["归属机构名称"] or org_names.get(values["归属机构代码"], "")
        values["人员状态"] = "正常"
        values["离职日期"] = ""
        if target == len(master):
            master.loc[target, MASTER_COLUMNS] = [values.get(column, "") for column in MASTER_COLUMNS]
        else:
            for column, value in values.items():
                master.loc[target, column] = value
        return
    if len(matches) != 1:
        issues.append({"issue_type": "change_not_matched", "message": f"{name} 的人员变更无法唯一匹配人员主数据"})
        return
    target = matches[0]
    if change_type == "离职":
        master.loc[target, "人员状态"] = "非正常"
        master.loc[target, "离职日期"] = _clean(row["离职日期"])
    elif change_type == "资料修改":
        for column in MASTER_COLUMNS:
            if column in {"姓名", "人员状态", "离职日期"}:
                continue
            if _clean(row[column]):
                master.loc[target, column] = _clean(row[column])
        if _clean(row["证件类型"]) == "居民身份证" and not _valid_id(_clean(row["证件号码"])):
            issues.append({"issue_type": "invalid_id_number", "message": f"{name} 的居民身份证号码校验不通过"})


def process_part_time(
    db,
    payroll: pd.DataFrame,
    changes: pd.DataFrame | None,
    *,
    period_id: int,
    stage: str = "initial",
) -> dict[str, Any]:
    period = db.query(Period).filter(Period.id == period_id).first()
    master, master_source = _load_previous_master_source(db, period_id)
    payroll = _normalize_payroll(payroll)
    changes = _normalize_changes(changes if changes is not None else pd.DataFrame(columns=CHANGE_COLUMNS))
    issues: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    if master_source["missing"]:
        expected = master_source["period"] or "上月"
        issues.append({
            "issue_type": "missing_previous_master",
            "message": f"未找到 {expected} 非全日制人员主数据，无法按上月基准核对，本次申报被阻断。",
        })
    if master_source["month_warning"]:
        warnings.append({
            "issue_type": "master_month_mismatch",
            "message": master_source["month_warning"],
            "severity": "warning",
        })
    if stage == "recheck" and changes.empty:
        issues.append({"issue_type": "missing_personnel_changes", "message": "请上传人员信息变更表后重新执行。"})
    if payroll.empty:
        issues.append({"issue_type": "missing_payroll", "message": "缺少非全日制用工工资表"})
    for index, row in payroll.iterrows():
        name = normalize_part_time_name(row["姓名"])
        if not name:
            issues.append({"issue_type": "missing_name", "message": f"工资表第 {index + 5} 行姓名为空"})
        if _number(row["本期收入"]) <= 0:
            issues.append({"issue_type": "invalid_income", "message": f"{name or '未命名人员'} 本期收入必须大于0"})
    payroll["_name"] = payroll["姓名"].map(normalize_part_time_name)
    # A name alone is not a unique person.  Same-name employees with distinct IDs
    # are valid; only repeated IDs (or repeated name rows without IDs) are duplicates.
    duplicate_keys = payroll.apply(
        lambda row: f"id:{row['证件号码']}" if _clean(row["证件号码"]) else f"name:{row['_name']}",
        axis=1,
    )
    for key in duplicate_keys[duplicate_keys.duplicated(keep=False)].unique():
        name = payroll.loc[duplicate_keys == key, "_name"].iloc[0]
        issues.append({"issue_type": "duplicate_payroll", "message": f"{name} 该人员本期存在多条工资记录，请在上传前合并为一条。"})

    org_names = _org_names(db)
    if stage == "recheck":
        for _, row in changes.iterrows():
            if _clean(row["姓名"]):
                _apply_change(master, row, org_names, issues)

    matched_rows: list[dict[str, Any]] = []
    matched_payroll_keys: set[tuple[str, str]] = set()
    matched_payroll_names: set[str] = set()
    seen_master: set[int] = set()
    pending: list[dict[str, Any]] = []
    automatic_rehires: list[dict[str, str]] = []
    for _, row in payroll.iterrows():
        name = row["_name"]
        candidates = master.index[master["姓名"].map(normalize_part_time_name) == name]
        active = [index for index in candidates if _clean(master.loc[index, "人员状态"]) == "正常"]
        historical = list(candidates)
        id_number = _clean(row["证件号码"])
        if len(active) > 1:
            if not id_number:
                issues.append({"issue_type": "same_name_missing_id", "message": f"{name}存在同名人员，请在工资表填写证件号码。"})
                continue
            active = [index for index in active if _clean(master.loc[index, "证件号码"]) == id_number]
        if len(active) == 0 and historical:
            rehire = [index for index in historical if _clean(master.loc[index, "人员状态"]) != "正常"]
            if len(rehire) == 1:
                if stage == "initial":
                    pending.append({"变更类型": "重新任职", "姓名": name, "证件号码": id_number})
                    issues.append({"issue_type": "rehire_requires_change", "message": f"{name} 为历史非正常人员，请在人员信息变更表补充重新任职日期。"})
                    continue
                active = rehire
                automatic_rehires.append({"姓名": name, "变化类型": "重新任职"})
        if len(active) != 1:
            pending.append({"变更类型": "新增", "姓名": name, "证件号码": id_number})
            issues.append({"issue_type": "new_part_time_person", "message": f"发现新增人员 {name}，请补充人员信息变更表"})
            continue
        index = active[0]
        seen_master.add(index)
        org_code = _clean(master.loc[index, "归属机构代码"])
        if not org_code:
            issues.append({"issue_type": "missing_org_code", "message": f"{name} 人员主数据缺少机构代码"})
            continue
        if org_code not in org_names:
            issues.append({"issue_type": "unknown_org_code", "message": f"{name} 的机构代码 {org_code} 不存在"})
            continue
        master_id = _clean(master.loc[index, "证件号码"])
        if id_number and id_number != master_id:
            warnings.append({"issue_type": "id_mismatch", "message": f"{name}工资表证件号码与人员主数据不一致，本次申报使用人员主数据证件号码。", "severity": "warning"})
        matched_rows.append({
            "工号": _clean(master.loc[index, "工号"]), "*姓名": master.loc[index, "姓名"],
            "*证件类型": master.loc[index, "证件类型"], "*证件号码": master_id,
            "*所得项目": PART_TIME_INCOME_ITEM, "本期收入": _number(row["本期收入"]),
            "本期免税收入": _number(row["本期免税收入"]), "商业健康保险": _number(row["商业健康保险"]),
            "税延养老保险": _number(row["税延养老保险"]), "其他": _number(row["其他"]),
            "允许扣除的税费": _number(row["允许扣除的税费"]), "减免税额": _number(row["减免税额"]),
            "协定减免": _number(row["协定减免"]), "备注": _clean(row["备注"]),
            "归属机构代码": org_code, "归属机构名称": _clean(master.loc[index, "归属机构名称"]) or org_names.get(org_code, ""),
        })
        matched_payroll_keys.add((name, id_number))
        matched_payroll_names.add(name)

    if period is not None:
        leave_date = (date(period.year, period.month, 1) - timedelta(days=1)).isoformat()
        for index, row in master.iterrows():
            if _clean(row["人员状态"]) == "正常" and index not in seen_master:
                master.loc[index, "人员状态"] = "非正常"
                master.loc[index, "离职日期"] = leave_date

    matching = []
    for row in payroll.to_dict("records"):
        name = normalize_part_time_name(row.get("姓名"))
        id_number = _clean(row.get("证件号码"))
        matched = (name, id_number) in matched_payroll_keys if id_number else name in matched_payroll_names
        matching.append({
            "工资表姓名": row.get("姓名", ""),
            "工资表证件号码": row.get("证件号码", ""),
            "匹配状态": "已匹配" if matched else "待补充",
        })
    changes_log = [{"姓名": row["姓名"], "变化类型": row["变更类型"]} for _, row in changes.iterrows() if _clean(row["姓名"])] if stage == "recheck" else []
    changes_log.extend(automatic_rehires)
    changes_log.extend({"姓名": row["姓名"], "变化类型": "自动离职"} for _, row in master.iterrows() if _clean(row["人员状态"]) == "非正常")
    summary = {
        "part_time_payroll_count": int(len(payroll)), "part_time_matched_count": int(len(matched_rows)),
        "part_time_new_count": int(sum(item["变更类型"] == "新增" for item in pending)),
        "part_time_rehire_count": int(sum(item["变更类型"] == "重新任职" for item in pending)),
        "part_time_leaver_count": int(sum(item["变化类型"] == "自动离职" for item in changes_log)),
        "part_time_change_count": int(len(changes)), "part_time_income_total": float(sum(item["本期收入"] for item in matched_rows)),
        "part_time_org_count": int(len({item["归属机构代码"] for item in matched_rows if item["归属机构代码"]})),
        "part_time_modified_count": int(sum(item.get("变化类型") == "资料修改" for item in changes_log)),
        "part_time_same_name_conflict_count": int(sum(item.get("issue_type") == "same_name_missing_id" for item in issues)),
        "payroll_count": int(len(payroll)), "matched_count": int(len(matched_rows)),
        "new_count": int(sum(item["变更类型"] == "新增" for item in pending)),
        "rehire_count": int(sum(item["变更类型"] == "重新任职" for item in pending)),
        "leaver_count": int(sum(item["变化类型"] == "自动离职" for item in changes_log)),
        "change_count": int(len(changes)), "income_total": float(sum(item["本期收入"] for item in matched_rows)),
        "org_count": int(len({item["归属机构代码"] for item in matched_rows if item["归属机构代码"]})),
        "modified_count": int(sum(item.get("变化类型") == "资料修改" for item in changes_log)),
        "same_name_conflict_count": int(sum(item.get("issue_type") == "same_name_missing_id" for item in issues)),
        "master_source_period": master_source["period"],
        "master_source_period_label": (
            f"{master_source['period'][:4]}年{master_source['period'][4:]}月"
            if len(master_source["period"]) == 6 else ""
        ),
        "master_source_file": master_source["file"],
        "master_source_materialized": False,
        "master_source_warning": master_source["month_warning"],
    }
    return {
        "master": master,
        "rows": pd.DataFrame(matched_rows, columns=DECLARATION_COLUMNS + ["归属机构代码", "归属机构名称"]),
        "issues": [*issues, *warnings],
        "blocking": issues,
        "pending": pd.DataFrame(pending, columns=CHANGE_COLUMNS),
        "matching": pd.DataFrame(matching),
        "changes": pd.DataFrame(changes_log),
        "summary": summary,
        "stage": stage,
        "master_source": master_source,
    }


def _write_template_rows(source_sheet: str, rows: pd.DataFrame, output: Path, header_override: list[str] | None = None) -> None:
    source_path = SOURCE_TEMPLATE_PATHS.get(source_sheet, TEMPLATE_PATH)
    workbook = load_workbook(source_path)
    sheet = workbook.active if source_sheet in SOURCE_TEMPLATE_PATHS else workbook[source_sheet]
    header_row = 1 if source_sheet in SOURCE_TEMPLATE_PATHS else 4
    data_start_row = header_row + 1
    headers = header_override or [sheet.cell(header_row, column).value for column in range(1, sheet.max_column + 1)]
    if header_override:
        for column, header in enumerate(header_override, start=1):
            sheet.cell(header_row, column).value = header
    for row_index in range(data_start_row, sheet.max_row + 1):
        for column in range(1, sheet.max_column + 1):
            sheet.cell(row_index, column).value = None
    for row_index, record in enumerate(rows.to_dict("records"), start=data_start_row):
        for column, header in enumerate(headers, start=1):
            key = header
            if key not in record and isinstance(key, str) and key.startswith("*"):
                key = key[1:]
            key = {
                "变更/调整类型": "变更类型",
                "证件号码（同名必填）": "*证件号码",
                "证件号码(同名必填)": "*证件号码",
                "*所得项目": "*所得项目",
            }.get(key, key)
            if key in record:
                sheet.cell(row_index, column).value = record[key]
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def write_part_time_artifacts(job_id: int, period: Period, result: dict[str, Any], *, finalized: bool) -> list[tuple[str, str]]:
    root = artifact_path(job_id, "part_time").parent
    artifacts: list[tuple[str, str]] = []
    if finalized:
        master_path = root / f"非全日制用工人员信息表_{period.year}{period.month:02d}.xlsx"
        _write_template_rows("人员信息表", result["master"], master_path)
        artifacts.append(("part_time_master", str(master_path)))
    if not result["pending"].empty:
        pending_path = root / f"非全日制人员信息变更表_待补充_{period.year}{period.month:02d}.xlsx"
        _write_template_rows("人员信息变更表", result["pending"].to_dict("records") and result["pending"], pending_path)
        artifacts.append(("part_time_pending_changes", str(pending_path)))
    if finalized and not result["blocking"]:
        for org, group in result["rows"].groupby("归属机构代码"):
            if not _clean(org):
                continue
            declaration = group.drop(columns=["归属机构代码", "归属机构名称"], errors="ignore")
            declaration_path = root / f"{org}_劳务报酬所得_{period.year}{period.month:02d}.xlsx"
            _write_template_rows("非全日制用工工资表", declaration, declaration_path, DECLARATION_COLUMNS)
            artifacts.append(("declaration", str(declaration_path)))
        changed = result["changes"]
        if not changed.empty:
            collection = result["master"][result["master"]["姓名"].isin(changed["姓名"].tolist())].copy()
            collection = collection.rename(columns={"姓名": "*姓名", "证件类型": "*证件类型", "证件号码": "*证件号码", "国籍（地区）": "*国籍(地区)", "性别": "*性别", "出生日期": "*出生日期", "任职/从业日期": "任职受雇从业日期"})
            for org, group in collection.groupby("归属机构代码", dropna=False):
                if not _clean(org):
                    continue
                collection_path = root / f"{org}_非全日制_人员信息采集导入_{period.year}{period.month:02d}.xlsx"
                _write_template_rows("人员信息表", group.to_dict("records") and group, collection_path)
                artifacts.append(("personnel_collection", str(collection_path)))
    workpaper = root / f"非全日制用工处理底稿_{period.year}{period.month:02d}.xlsx"
    write_workbook(workpaper, {"申报明细": result["rows"], "人员匹配结果": result["matching"], "人员变化": result["changes"], "问题清单": pd.DataFrame(result["issues"]), "机构汇总": result["rows"].groupby(["归属机构代码", "归属机构名称"], dropna=False).agg(人数=("*姓名", "count"), 收入合计=("本期收入", "sum"), 免税收入=("本期免税收入", "sum"), 其他扣除=("允许扣除的税费", "sum")).reset_index()})
    artifacts.append(("part_time_workpaper", str(workpaper)))
    return artifacts


def save_part_time_master(db, period_id: int, master: pd.DataFrame) -> PersonnelMasterArtifact:
    period = db.query(Period).filter(Period.id == period_id).first()
    root = settings.artifact_dir / "personnel_masters" / str(period_id) / "part_time"
    root.mkdir(parents=True, exist_ok=True)
    path = root / f"{period.year}{period.month:02d}_part_time_month_all.xlsx"
    write_workbook(path, {"人员信息表": master}, text_values=True)
    artifact = db.query(PersonnelMasterArtifact).filter(PersonnelMasterArtifact.period_id == period_id, PersonnelMasterArtifact.person_type == "part_time", PersonnelMasterArtifact.scope_type == "month", PersonnelMasterArtifact.scope_code == "").first()
    if artifact is None:
        artifact = PersonnelMasterArtifact(period_id=period_id, person_type="part_time", scope_type="month", scope_code="", file_name=path.name, stored_path=str(path), row_count=len(master), validation_issues=[])
        db.add(artifact)
    else:
        artifact.file_name, artifact.stored_path, artifact.row_count = path.name, str(path), len(master)
    artifact.updated_at = datetime.utcnow()
    db.flush()
    db.add(PersonnelMasterImportBatch(
        period_id=period_id,
        artifact_id=artifact.id,
        person_type="part_time",
        scope_type="month",
        scope_code="",
        original_name="非全日制用工流程自动生成本月人员主数据",
        stored_path=str(path),
        row_count=len(master),
        validation_issues=[],
    ))
    return artifact
