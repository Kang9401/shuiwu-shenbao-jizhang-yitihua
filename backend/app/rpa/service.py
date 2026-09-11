from __future__ import annotations

import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import threading
import time
import urllib.request
import uuid
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models.core import Artifact, Job
from app.rpa.commands import build_chrome_command, build_task_command
from app.rpa.log_parser import parse_log_line
from app.rpa.paths import company_rpa_root, state_dir, uploads_dir
from app.rpa.process_manager import ProcessManager
from app.rpa.popup_rules import default_popup_rules, normalize_popup_rules
from app.rpa.runtime import ensure_runtime_app, sha256_file
from app.rpa.state_store import StateStore, result_cell
from app.services.personnel_master import list_rpa_organizations

TASK_NAMES = {
    "special_deduction": "专项附加导出",
    "import": "导入数据",
    "tax_certificate": "完税证明下载",
    "income_report": "综合所得申报表下载",
    "extra_income_reports": "分类/限售股申报表下载",
    "declaration_reports": "申报结果下载",
}
NAME_HEADERS = {"机构名称", "机构简称", "单位名称", "名称"}
CODE_HEADERS = {"机构代码", "机构编号", "代码", "编号"}


def now_text() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def safe_name(name: str) -> str:
    if not name or name != Path(name).name or name in {".", ".."} or ".." in Path(name).parts:
        raise ValueError("文件名不安全")
    return name


def inside(root: Path, candidate: Path) -> Path:
    resolved_root = root.resolve()
    resolved = candidate.resolve()
    if resolved != resolved_root and resolved_root not in resolved.parents:
        raise ValueError("文件路径超出允许目录")
    return resolved


class RpaService:
    def __init__(self, store: StateStore | None = None, manager: ProcessManager | None = None):
        self.store = store or StateStore()
        self.manager = manager or ProcessManager()
        self.company_id: int | None = None
        self._custom_store = store is not None
        self._company_lock = threading.RLock()
        self._log_lock = threading.RLock()
        self._cancel_requested = False
        self._stderr_lines: list[str] = []
        self._composite: dict | None = None

    def initialize(self) -> None:
        ensure_runtime_app()
        self.store.recover_interrupted()

    def shutdown(self) -> None:
        self.stop_task()

    def activate_company(self, company_id: int) -> None:
        with self._company_lock:
            if self.company_id == company_id:
                return
            if self.manager.running:
                raise RuntimeError("RPA 任务运行中，不能切换分公司")
            self.company_id = company_id
            if not self._custom_store:
                self.store = StateStore(state_dir(company_id) / "rpa_state.json")
                self.store.recover_interrupted()
            self._write_runtime_config()

    def is_company_running(self, company_id: int) -> bool:
        return self.manager.running and self.company_id == company_id

    def _state_dir(self) -> Path:
        return state_dir(self.company_id)

    def _configured_root(self, key: str, fallback: str) -> Path:
        state = self.store.read() if hasattr(self.store, "read") else {}
        value = str(state.get("config", {}).get(key) or "").strip()
        if value:
            root = Path(value).expanduser()
        elif self.company_id is not None:
            root = company_rpa_root(self.company_id) / fallback
        else:
            root = ensure_runtime_app() / fallback
        root.mkdir(parents=True, exist_ok=True)
        return root.resolve()

    def _input_root(self) -> Path:
        return self._configured_root("input_path", "input")

    def _output_root(self) -> Path:
        return self._configured_root("output_path", "output")

    def _write_runtime_config(self) -> None:
        data = self.store.read()
        stored_rules = data.get("popup_rules")
        config = {
            "chrome_path": str(data.get("config", {}).get("chrome_path") or ""),
            "input_path": str(self._input_root()),
            "output_path": str(self._output_root()),
            "popup_rules": normalize_popup_rules(
                default_popup_rules() if stored_rules is None else stored_rules,
                drop_protected=True,
            ),
        }
        path = ensure_runtime_app() / "etax_config.json"
        temp = path.with_suffix(".tmp")
        temp.write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")
        temp.replace(path)

    def _clear_runtime_output(self, preserve_names: set[str] | None = None) -> None:
        root = (ensure_runtime_app() / "output").resolve()
        root.mkdir(parents=True, exist_ok=True)
        preserve_names = preserve_names or set()
        for path in root.iterdir():
            if path.is_file() and path.name in preserve_names:
                continue
            if path.is_dir():
                shutil.rmtree(path)
            else:
                path.unlink(missing_ok=True)

    def _restore_resume_reconciliation(self, month: str) -> set[str]:
        """Bring a prior configured-output reconciliation workbook back to runtime."""
        name = f"{month}个税申报核对.xlsx"
        source = inside(self._output_root(), self._output_root() / name)
        target_root = (ensure_runtime_app() / "output").resolve()
        target = target_root / name
        if source.is_file() and source.resolve() != target.resolve():
            shutil.copy2(source, target)
        return {name} if target.is_file() else set()

    def _sync_output_files(self) -> None:
        source = (ensure_runtime_app() / "output").resolve()
        target = self._output_root()
        if source != target:
            for path in source.iterdir():
                if path.is_file():
                    shutil.copy2(path, target / path.name)

    def _read_orgs(self) -> list[dict[str, str]]:
        path = ensure_runtime_app() / "机构信息表.xlsx"
        if not path.is_file():
            raise ValueError("请先选择机构 Excel")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = [str(value or "").strip().replace(" ", "").replace("\u3000", "") for value in next(rows, [])]
            name_index = next((i for i, value in enumerate(headers) if value in NAME_HEADERS), None)
            code_index = next((i for i, value in enumerate(headers) if value in CODE_HEADERS), None)
            result_index_column = next((i for i, value in enumerate(headers) if value in {"RPA搜索结果序号", "RPA机构序号"}), None)
            if name_index is None or code_index is None:
                raise ValueError("机构信息表必须包含机构名称和机构代码列")
            result, seen = [], set()
            for row in rows:
                name = str(row[name_index] or "").strip() if name_index < len(row) else ""
                code = str(row[code_index] or "").strip() if code_index < len(row) else ""
                if code.endswith(".0"):
                    code = code[:-2]
                result_index = 1
                if result_index_column is not None and result_index_column < len(row) and row[result_index_column] not in (None, ""):
                    try:
                        result_index = int(float(str(row[result_index_column]).strip()))
                    except (TypeError, ValueError):
                        raise ValueError(f"机构 {code or name} 的 RPA 搜索结果序号不是有效正整数") from None
                    if result_index < 1:
                        raise ValueError(f"机构 {code or name} 的 RPA 搜索结果序号必须从 1 开始")
                if name and code and (name, code) not in seen:
                    seen.add((name, code))
                    result.append({"code": code, "name": name, "rpa_search_result_index": result_index})
            if not result:
                raise ValueError("机构信息表没有有效机构数据")
            return result
        finally:
            workbook.close()

    def preview_orgs(self, all_orgs: bool, org_code: str | None, start_org_code: str | None = None) -> list[dict[str, str]]:
        orgs = self._read_orgs()
        if start_org_code:
            index = next((i for i, item in enumerate(orgs) if item["code"] == start_org_code), None)
            if index is None:
                raise ValueError(f"机构信息表中找不到机构代码：{start_org_code}")
            return orgs[index:]
        if all_orgs:
            return orgs
        match = next((item for item in orgs if item["code"] == (org_code or "").strip()), None)
        if not match:
            raise ValueError(f"机构信息表中找不到机构代码：{org_code or ''}")
        return [match]

    def organizations_for_period(self, db: Session, period_id: int) -> list[dict]:
        return list_rpa_organizations(db, period_id=period_id)

    def prepare_runtime_org_excel(self, db: Session, *, period_id: int, selected_org_codes=None):
        organizations = self.organizations_for_period(db, period_id)
        requested = [str(code).strip() for code in (selected_org_codes or []) if str(code).strip()]
        by_code = {item["code"]: item for item in organizations}
        missing = [code for code in requested if code not in by_code]
        if missing:
            raise ValueError(f"所选机构不在当前期间人员主数据中：{'、'.join(missing)}")
        selected = [item for item in organizations if not requested or item["code"] in set(requested)]
        target = ensure_runtime_app() / "机构信息表.xlsx"
        temp = target.with_name(f".{target.stem}-{uuid.uuid4().hex}.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "机构信息"
        include_result_index = any(item.get("rpa_search_result_index", 1) != 1 for item in selected)
        sheet.append(["机构名称", "机构代码", "RPA搜索结果序号"] if include_result_index else ["机构名称", "机构代码"])
        for item in selected:
            values = [item["name"], item["code"]]
            if include_result_index:
                values.append(item.get("rpa_search_result_index", 1))
            sheet.append(values)
            sheet.cell(sheet.max_row, 2).number_format = "@"
        workbook.save(temp)
        temp.replace(target)
        reread = self._read_orgs()
        if [(item["code"], item["name"], item.get("rpa_search_result_index", 1)) for item in reread] != [(item["code"], item["name"], item.get("rpa_search_result_index", 1)) for item in selected]:
            raise RuntimeError("运行时机构信息表回读校验失败")
        self.store.update(lambda data: data["config"].update(org_excel_path=None, org_excel_name="人员主数据自动生成"))
        return target, selected

    def chrome_status(self) -> dict:
        status, message = "stopped", "未检测到可接管的 Chrome"
        try:
            with urllib.request.urlopen("http://127.0.0.1:9222/json/version", timeout=1) as response:
                payload = json.loads(response.read().decode("utf-8"))
            if payload.get("webSocketDebuggerUrl"):
                status, message = "ready", "CDP 已连接"
            else:
                status, message = "unavailable", "CDP 返回内容无效"
        except Exception:
            pass
        value = {"status": status, "last_checked_at": now_text(), "message": message}
        self.store.update(lambda data: data.__setitem__("chrome", value))
        return value

    def get_status(self, refresh_chrome: bool = True) -> dict:
        if refresh_chrome:
            self.chrome_status()
        data = self.store.read()
        results = []
        for month, orgs in data.get("results", {}).items():
            for code, values in orgs.items():
                for task_key in TASK_NAMES:
                    values.setdefault(task_key, "待处理")
                results.append({"month": month, "code": code, **values})
        data["config"]["input_path"] = str(self._input_root())
        data["config"]["output_path"] = str(self._output_root())
        import_files = self._file_items(self._input_root(), include_url=False)
        run_status = data["current_run"].get("status")
        data["results"] = results
        data["import_files"] = import_files
        data["can_start"] = run_status not in {"starting", "running", "stopping"} and data["chrome"].get("status") == "ready"
        data["can_stop"] = self.manager.running and run_status != "stopping"
        failure = data.get("last_failure", {})
        data["can_resume"] = not self.manager.running and bool(failure.get("task_key") and failure.get("org_code"))
        data["company_id"] = self.company_id
        return data

    def save_config(self, chrome_path: str, input_path: str = "", output_path: str = "") -> dict:
        if self.manager.running:
            raise RuntimeError("RPA 任务运行中，不能修改路径配置")
        value = chrome_path.strip()
        app_dir = ensure_runtime_app()
        input_root = Path(input_path.strip()).expanduser() if input_path.strip() else app_dir / "input"
        output_root = Path(output_path.strip()).expanduser() if output_path.strip() else app_dir / "output"
        if not input_root.is_absolute() or not output_root.is_absolute():
            raise ValueError("INPUT 和 OUTPUT 必须使用完整路径")
        input_root.mkdir(parents=True, exist_ok=True)
        output_root.mkdir(parents=True, exist_ok=True)
        if input_root.resolve() == output_root.resolve():
            raise ValueError("INPUT 和 OUTPUT 不能使用同一个目录")
        config = {"chrome_path": value, "input_path": str(input_root.resolve()), "output_path": str(output_root.resolve())}
        self.store.update(lambda data: data["config"].update(config))
        self._write_runtime_config()
        return self.store.read()["config"]

    def get_popup_rules(self) -> list[dict]:
        stored_rules = self.store.read().get("popup_rules")
        return normalize_popup_rules(
            default_popup_rules() if stored_rules is None else stored_rules,
            drop_protected=True,
        )

    def save_popup_rules(self, rules: list[dict]) -> list[dict]:
        if self.manager.running:
            raise RuntimeError("RPA 任务运行中，不能修改弹窗规则")
        normalized = normalize_popup_rules(rules)
        self.store.update(lambda data: data.__setitem__("popup_rules", normalized))
        self._write_runtime_config()
        return normalized

    def reset_popup_rules(self) -> list[dict]:
        return self.save_popup_rules(default_popup_rules())

    def popup_events(self, limit: int = 100) -> list[dict]:
        paths = [ensure_runtime_app() / "output" / "popup_events.jsonl"]
        configured = self._output_root() / "popup_events.jsonl"
        if configured.resolve() != paths[0].resolve():
            paths.append(configured)
        events: list[dict] = []
        seen: set[str] = set()
        for path in paths:
            if not path.is_file():
                continue
            try:
                lines = path.read_text(encoding="utf-8").splitlines()
            except OSError:
                continue
            for line in lines[-limit:]:
                try:
                    event = json.loads(line)
                except json.JSONDecodeError:
                    continue
                identity = json.dumps(event, ensure_ascii=False, sort_keys=True)
                if identity not in seen:
                    seen.add(identity)
                    events.append(event)
        events.sort(key=lambda item: str(item.get("captured_at") or ""), reverse=True)
        return events[:limit]

    def upload_org_excel(self, filename: str, content: bytes) -> dict:
        name = safe_name(filename)
        if Path(name).suffix.lower() != ".xlsx":
            raise ValueError("机构信息表只支持 .xlsx")
        upload = uploads_dir(self.company_id) / "org_excel" / f"{uuid.uuid4().hex}_{name}"
        upload.parent.mkdir(parents=True, exist_ok=True)
        upload.write_bytes(content)
        target = ensure_runtime_app() / "机构信息表.xlsx"
        shutil.copy2(upload, target)
        try:
            count = len(self._read_orgs())
        except Exception:
            target.unlink(missing_ok=True)
            upload.unlink(missing_ok=True)
            raise
        def change(data: dict) -> None:
            data["config"].update(org_excel_path=str(upload), org_excel_name=name)
        self.store.update(change)
        return {"name": name, "count": count}

    def upload_import_files(self, files: list[tuple[str, bytes]], overwrite: bool) -> list[dict]:
        root = self._input_root()
        prepared = []
        for filename, content in files:
            name = safe_name(filename)
            if Path(name).suffix.lower() not in {".xls", ".xlsx"}:
                raise ValueError(f"导入文件只支持 .xls/.xlsx：{name}")
            target = inside(root, root / name)
            if target.exists() and not overwrite:
                raise FileExistsError(name)
            prepared.append((target, content))
        for target, content in prepared:
            target.write_bytes(content)
        return self._file_items(root, include_url=False)

    def prepare_from_period(self, db: Session, period_id: int, overwrite: bool) -> list[dict]:
        jobs = db.query(Job).filter(Job.period_id == period_id, Job.status == "success", Job.operation == "generate").order_by(Job.created_at.desc()).all()
        latest: dict[str, Job] = {}
        for job in jobs:
            latest.setdefault(job.workflow_code, job)
        artifacts = db.query(Artifact).filter(Artifact.job_id.in_([job.id for job in latest.values()]), Artifact.artifact_type.in_(["declaration", "personnel_collection"])).all() if latest else []
        candidates: dict[str, Path] = {}
        for artifact in artifacts:
            name = safe_name(artifact.file_name)
            source = Path(artifact.stored_path).resolve()
            if not source.is_file():
                continue
            prior = candidates.get(name)
            if prior and sha256_file(prior) != sha256_file(source):
                raise FileExistsError(f"本期产物存在同名不同内容文件：{name}")
            candidates[name] = source
        if not candidates:
            raise ValueError("本期没有可用于 RPA 导入的已生成文件")
        root = self._input_root()
        for name, source in candidates.items():
            target = inside(root, root / name)
            if target.exists() and sha256_file(target) == sha256_file(source):
                continue
            if target.exists() and not overwrite:
                raise FileExistsError(name)
        for name, source in candidates.items():
            target = root / name
            if not target.exists() or sha256_file(target) != sha256_file(source):
                shutil.copy2(source, target)
        return self._file_items(root, include_url=False)

    def delete_import_files(self, names: list[str]) -> list[dict]:
        root = self._input_root()
        for raw in names:
            name = safe_name(raw)
            target = inside(root, root / name)
            if not target.is_file():
                raise ValueError(f"导入文件不存在：{name}")
        for raw in names:
            (root / raw).unlink()
        return self._file_items(root, include_url=False)

    def start_chrome(self, chrome_path: str = "") -> dict:
        if self.chrome_status()["status"] == "ready":
            return {"message": "可接管的 Chrome 已在运行"}
        app_dir = ensure_runtime_app()
        command = build_chrome_command(app_dir, chrome_path)
        (app_dir / ".chrome-debug-profile").mkdir(parents=True, exist_ok=True)
        try:
            process = subprocess.Popen(command, cwd=str(app_dir), shell=False)
        except OSError as exc:
            raise RuntimeError(f"Chrome 启动失败：{exc}") from exc
        self.store.update(lambda data: data["chrome"].update(status="starting", message="正在启动可接管的 Chrome"))
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            status = self.chrome_status()
            if status["status"] == "ready":
                return {"message": "已启动可接管的 Chrome。请在该窗口手工登录自然人电子税务局，登录完成后回到本页面执行任务。"}
            time.sleep(0.25)
        exit_code = process.poll()
        detail = f"Chrome 进程退出码 {exit_code}" if exit_code is not None else "Chrome 已启动但调试端口 9222 未就绪"
        self.store.update(lambda data: data["chrome"].update(status="unavailable", message=detail))
        raise RuntimeError(f"Chrome 初始化失败：{detail}。请关闭已有的 RPA Chrome 后重试。")

    def start_task(self, task_key: str, month: str, all_orgs: bool, org_code: str | None, resume_mode: str | None = None, start_org_code: str | None = None) -> dict:
        if self.manager.running:
            raise RuntimeError("已有 RPA 任务正在运行")
        if self.chrome_status()["status"] != "ready":
            raise RuntimeError("可接管的 Chrome 未就绪，请先初始化并手工登录")
        app_dir = ensure_runtime_app()
        orgs = self.preview_orgs(all_orgs, org_code, start_org_code)
        input_root = self._input_root()
        if task_key == "import" and not any(input_root.iterdir()):
            raise ValueError("input 中没有可导入文件")
        self._write_runtime_config()
        preserve_names: set[str] = set()
        if task_key == "import" and resume_mode == "resume":
            preserve_names = self._restore_resume_reconciliation(month)
        self._clear_runtime_output(preserve_names)
        resume_output_preserved = bool(preserve_names)
        actual_task = "income_report" if task_key == "declaration_reports" else task_key
        command, backend_month = build_task_command(app_dir, actual_task, month, all_orgs, org_code, resume_mode, start_org_code, input_root=input_root)
        run_id, started = uuid.uuid4().hex, now_text()
        log_dir = self._state_dir() / "logs"
        log_dir.mkdir(parents=True, exist_ok=True)
        log_path = log_dir / f"{datetime.now():%Y%m%d_%H%M%S}_{task_key}_{run_id[:8]}.log"
        log_path.write_text("", encoding="utf-8")
        for old_log in sorted(log_dir.glob("*.log"), key=lambda item: item.stat().st_mtime, reverse=True)[30:]:
            old_log.unlink(missing_ok=True)
        def starting(data: dict) -> None:
            month_results = data["results"].setdefault(month, {})
            for item in orgs:
                row = month_results.setdefault(item["code"], {"name": item["name"], **{key: "待处理" for key in TASK_NAMES if key != "declaration_reports"}})
                row["name"] = item["name"]
                for key in TASK_NAMES:
                    row.setdefault(key, "待处理")
                if task_key == "declaration_reports":
                    row["comprehensive_income_report"] = result_cell()
                    row["classified_income_report"] = result_cell()
                    row["restricted_stock_report"] = result_cell()
                else:
                    row[task_key] = "待处理"
            data["current_run"].update(run_id=run_id, task_key=task_key, subtask_key=actual_task if task_key == "declaration_reports" else None, subtask_index=1 if task_key == "declaration_reports" else None, subtask_count=2 if task_key == "declaration_reports" else None, display_name=TASK_NAMES[task_key], month=month, declaration_month=month, backend_month=backend_month, all_orgs=all_orgs, org_code=org_code, current_org_code=None, current_org_name=None, status="starting", pid=None, started_at=started, finished_at=None, exit_code=None, error_message=None, log_path=str(log_path))
        self.store.update(starting)
        if resume_output_preserved:
            # The import runner appends/replaces institution sheets in place.
            self._handle_line(f"保留已有核对表并追加机构：{month}个税申报核对.xlsx")
        self._cancel_requested = False
        self._stderr_lines = []
        self._composite = {"app_dir": app_dir, "month": month, "all_orgs": all_orgs, "org_code": org_code, "start_org_code": start_org_code} if task_key == "declaration_reports" else None
        try:
            pid = self.manager.start(command, app_dir, self._handle_line, self._handle_exit)
        except Exception as exc:
            self._mark_start_failure(str(exc))
            raise
        self.store.update(lambda data: data["current_run"].update(status="running", pid=pid))
        return {"run_id": run_id, "pid": pid, "message": "任务已启动"}

    def _mark_start_failure(self, message: str) -> None:
        def change(data: dict) -> None:
            run = data["current_run"]
            run.update(status="failed", finished_at=now_text(), error_message=message, pid=None)
        self.store.update(change)

    def _handle_line(self, line: str, is_stderr: bool = False) -> None:
        data = self.store.read()
        path = Path(data["current_run"].get("log_path") or self._state_dir() / "current.log")
        with self._log_lock, path.open("a", encoding="utf-8", newline="\n") as handle:
            handle.write(line + "\n")
        if is_stderr:
            self._stderr_lines = [*self._stderr_lines[-49:], line]
        task_key = data["current_run"].get("subtask_key") or data["current_run"].get("task_key")
        current = data["current_run"].get("current_org_code")
        events = parse_log_line(line, task_key, current)
        if not events:
            return
        def change(state: dict) -> None:
            run = state["current_run"]
            rows = state["results"].setdefault(run["month"], {})
            for event in events:
                if event.kind == "start":
                    run.update(current_org_code=event.org_code, current_org_name=event.org_name)
                    row = rows.setdefault(event.org_code, {"name": event.org_name, **{key: "待处理" for key in TASK_NAMES if key != "declaration_reports"}})
                    result_key = "comprehensive_income_report" if task_key == "income_report" and run.get("task_key") == "declaration_reports" else task_key
                    row[result_key] = result_cell("running") if run.get("task_key") == "declaration_reports" else "处理中"
                elif event.kind == "increment":
                    row = rows.get(event.org_code, {})
                    found = re.search(r"成功 (\d+)笔", row.get(task_key, ""))
                    row[task_key] = f"成功 {(int(found.group(1)) if found else 0) + 1}笔"
                elif event.kind == "done":
                    row = rows.get(event.org_code, {})
                    result_key = "comprehensive_income_report" if task_key == "income_report" and run.get("task_key") == "declaration_reports" else task_key
                    if run.get("task_key") == "declaration_reports" and event.count is not None:
                        row[result_key] = result_cell("skipped" if event.count == 0 else "success", count=event.count, reason="查询成功但无符合条件的记录" if event.count == 0 else None)
                    elif event.count is not None:
                        row[result_key] = f"成功 {event.count}笔"
                    elif not str(row.get(task_key, "")).startswith("成功"):
                        row[task_key] = "成功 0笔"
                elif event.kind == "result":
                    row = rows.get(event.org_code, {})
                    key = {"classified_income": "classified_income_report", "restricted_stock": "restricted_stock_report"}.get(event.category)
                    if key:
                        row[key] = result_cell(event.status or "failed", count=event.count, reason="查询成功但无符合条件的记录" if event.status == "skipped" else event.reason)
        self.store.update(change)

    def _handle_exit(self, exit_code: int, cancelled: bool) -> None:
        finished = now_text()
        data = self.store.read()
        run = data["current_run"]
        cancelled = cancelled or self._cancel_requested
        if self._composite and run.get("task_key") == "declaration_reports" and run.get("subtask_key") == "income_report" and exit_code == 0 and not cancelled:
            app_dir = self._composite["app_dir"]
            command, backend_month = build_task_command(app_dir, "extra_income_reports", self._composite["month"], self._composite["all_orgs"], self._composite["org_code"], None, self._composite["start_org_code"])
            self._handle_line("========== 开始子任务：分类所得与限售股申报表下载 ==========", False)
            self.store.update(lambda state: state["current_run"].update(subtask_key="extra_income_reports", subtask_index=2, backend_month=backend_month, current_org_code=None, current_org_name=None, pid=None))
            self._stderr_lines = []
            try:
                pid = self.manager.start(command, app_dir, self._handle_line, self._handle_exit)
                self.store.update(lambda state: state["current_run"].update(pid=pid, status="running"))
                return
            except Exception as exc:
                exit_code = 1
                self._stderr_lines = [str(exc)]
        def change(data: dict) -> None:
            run = data["current_run"]
            status = "cancelled" if cancelled else ("succeeded" if exit_code == 0 else "failed")
            if status == "failed" and run.get("current_org_code"):
                failure_key = run.get("subtask_key") or run["task_key"]
                result_key = {"income_report": "comprehensive_income_report", "extra_income_reports": "classified_income_report"}.get(failure_key, failure_key)
                current_value = data["results"][run["month"]][run["current_org_code"]].get(result_key)
                data["results"][run["month"]][run["current_org_code"]][result_key] = result_cell("failed", error_message="RPA 子进程异常退出") if isinstance(current_value, dict) else "失败"
                if failure_key == "extra_income_reports":
                    data["results"][run["month"]][run["current_org_code"]]["restricted_stock_report"] = result_cell("failed", error_message="RPA 子进程异常退出，无法判断结果")
                data["last_failure"] = {"task_key": run["task_key"], "org_code": run["current_org_code"], "month": run["month"]}
            error_message = None if exit_code == 0 else ("\n".join(self._stderr_lines).strip() or f"RPA 进程退出码：{exit_code}")
            run.update(status=status, finished_at=finished, exit_code=exit_code, pid=None, error_message=error_message)
            data["history"].insert(0, {"task_key": run["task_key"], "task": run["display_name"], "month": run["month"], "started_at": run["started_at"], "finished_at": finished, "status": "已取消" if cancelled else ("成功" if exit_code == 0 else "失败")})
        self.store.update(change)
        self._sync_output_files()
        self._composite = None

    def stop_task(self) -> dict:
        run_status = self.store.read()["current_run"].get("status")
        if not self.manager.running and run_status not in {"starting", "running", "stopping"}:
            return {"message": "当前没有正在运行的任务"}
        self._cancel_requested = True
        self.store.update(lambda data: data["current_run"].update(status="stopping", error_message=None))
        stopped = self.manager.stop()
        deadline = time.monotonic() + 3
        while time.monotonic() < deadline:
            state = self.store.read()["current_run"]
            if state.get("status") != "stopping" and not self.manager.running:
                return {"message": "RPA 任务及相关子进程已停止"}
            time.sleep(0.05)
        if not self.manager.running:
            def finish_cancel(data: dict) -> None:
                run = data["current_run"]
                if run.get("status") == "stopping":
                    run.update(status="cancelled", finished_at=now_text(), pid=None, exit_code=None, error_message=None)
            self.store.update(finish_cancel)
            return {"message": "RPA 任务及相关子进程已停止"}
        raise RuntimeError("RPA 停止请求未能确认全部相关进程退出")

    def reset_results(self, month: str | None = None) -> dict:
        if self.manager.running:
            raise RuntimeError("RPA 任务运行中，不能重置处理结果")

        def change(data: dict) -> None:
            if month:
                data.get("results", {}).pop(month, None)
                if data.get("last_failure", {}).get("month") == month:
                    data["last_failure"] = {"task_key": None, "org_code": None, "month": None}
            else:
                data["results"] = {}
                data["last_failure"] = {"task_key": None, "org_code": None, "month": None}

        self.store.update(change)
        return self.get_status(refresh_chrome=False)

    def resume_task(self) -> dict:
        failure = self.store.read().get("last_failure", {})
        if not failure.get("task_key") or not failure.get("org_code") or not failure.get("month"):
            raise ValueError("当前没有可续跑的失败机构")
        return self.start_task(failure["task_key"], failure["month"], True, None, "resume" if failure["task_key"] == "import" else None, failure["org_code"])

    def read_logs(self, offset: int) -> dict:
        data = self.store.read()
        configured = data.get("current_run", {}).get("log_path")
        path = Path(configured) if configured else self._state_dir() / "current.log"
        if not path.is_file():
            return {"text": "", "next_offset": 0, "finished": not self.manager.running}
        size = path.stat().st_size
        start = max(0, min(offset, size))
        with path.open("rb") as handle:
            handle.seek(start)
            raw = handle.read()
        return {"text": raw.decode("utf-8", errors="replace"), "next_offset": start + len(raw), "finished": not self.manager.running}

    @staticmethod
    def _file_items(root: Path, include_url: bool) -> list[dict]:
        items = []
        if not root.is_dir():
            return items
        for path in sorted((item for item in root.iterdir() if item.is_file()), key=lambda item: item.stat().st_mtime, reverse=True):
            item = {"name": path.name, "size": path.stat().st_size, "modified_at": datetime.fromtimestamp(path.stat().st_mtime).astimezone().isoformat(timespec="seconds")}
            if include_url:
                item["download_url"] = f"/api/rpa/files/download?name={path.name}"
            items.append(item)
        return items

    def list_output_files(self) -> list[dict]:
        return self._file_items(self._output_root(), include_url=True)

    def resolve_output(self, name: str) -> Path:
        root = self._output_root()
        target = inside(root, root / safe_name(name))
        if not target.is_file():
            raise ValueError("输出文件不存在")
        return target

    def build_output_archive(self) -> tuple[str, bytes]:
        if self.manager.running:
            raise RuntimeError("RPA 任务运行中，不能打包输出文件")
        root = self._output_root()
        files = sorted(
            (path for path in root.rglob("*") if path.is_file() and not path.name.startswith("~$")),
            key=lambda item: item.relative_to(root).as_posix(),
        )
        if not files:
            raise ValueError("当前没有可下载的输出文件")
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
            for path in files:
                archive.write(path, arcname=path.relative_to(root).as_posix())
        content = buffer.getvalue()
        with zipfile.ZipFile(io.BytesIO(content), "r") as archive:
            bad_file = archive.testzip()
            if bad_file:
                raise RuntimeError(f"ZIP 完整性校验失败：{bad_file}")
            if len(archive.infolist()) != len(files):
                raise RuntimeError("ZIP 文件数量校验失败")
        return f"RPA输出文件_{datetime.now():%Y%m%d_%H%M%S}.zip", content

    def clear_output_files(self) -> list[dict]:
        if self.manager.running:
            raise RuntimeError("RPA 任务运行中，不能清空输出文件")
        root = self._output_root()
        for path in root.iterdir():
            if path.is_file():
                inside(root, path).unlink()
        runtime_root = (ensure_runtime_app() / "output").resolve()
        if runtime_root != root.resolve():
            for path in runtime_root.iterdir():
                if path.is_file():
                    inside(runtime_root, path).unlink()
        return self.list_output_files()


rpa_service = RpaService()
