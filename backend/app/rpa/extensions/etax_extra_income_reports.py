from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import shutil
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

try:
    from .etax_report_download import capture_sanitized_failure, check_cancelled, interruptible_wait, run_report_download
except ImportError:  # Runtime copies extension modules into one directory.
    from etax_report_download import capture_sanitized_failure, check_cancelled, interruptible_wait, run_report_download


WORK_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
OUTPUT_DIR = WORK_DIR / "output"
CONFIG_FILE = WORK_DIR / "etax_extra_income_reports_config.json"
CONFIG_ENV = "ETAX_EXTRA_INCOME_REPORTS_CONFIG"


@dataclass(frozen=True)
class ReportSpec:
    key: str
    title: str
    route: str
    export_label: str
    keyword: str


REPORT_SPECS = (
    ReportSpec(
        "classified_income",
        "分类所得申报表",
        "#/withholding/classified_income_declaration",
        "个人所得税扣缴申报表",
        "分类所得",
    ),
    ReportSpec(
        "restricted_stock",
        "限售股所得申报表",
        "#/withholding/restricted_shares",
        "限售股转让所得个人所得税扣缴报告表",
        "限售股",
    ),
)

# Captured read-only from the authorized tax site on 2026-07-22. A runtime JSON
# or environment value can override these values if the site DOM changes.
SITE_CONFIG_DEFAULTS = {
    "current_org_scope": ".company-name",
    "month_input": ".tax-period-date-picker input",
    "page_scope": ".main-page",
    "status_value": ".declare-status .dstatus",
    "success_record": ".declare-success",
    "export_button": ".dropdown-export button",
    "export_menu_item": ".el-dropdown-menu:visible .el-dropdown-menu__item",
    "export_dialog": ".export-result-list-message-dialog:visible",
    "export_rows": ".el-table__body .el-table__row",
    "export_refresh": ".modal-header-refresh",
    "export_download": "a",
}


def log(message: str) -> None:
    print(f"[{dt.datetime.now():%Y-%m-%d %H:%M:%S}] {message}", flush=True)


def load_site_config(path: Path = CONFIG_FILE) -> dict[str, str]:
    values: dict[str, Any] = dict(SITE_CONFIG_DEFAULTS)
    if path.is_file():
        values.update(json.loads(path.read_text(encoding="utf-8")))
    raw = os.environ.get(CONFIG_ENV, "").strip()
    if raw:
        values.update(json.loads(raw))
    missing = [key for key in SITE_CONFIG_DEFAULTS if not isinstance(values.get(key), str) or not values[key].strip()]
    if missing:
        raise RuntimeError(
            "分类/限售股申报表下载尚未配置已确认的税务网站选择器："
            + "、".join(missing)
            + f"。请将授权页面采集值写入 {path.name} 后再执行。"
        )
    return {key: str(values[key]).strip() for key in SITE_CONFIG_DEFAULTS}


def safe_filename_part(value: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]+', "_", str(value).strip())
    return cleaned.strip(" .") or "未命名"


def report_path(spec: ReportSpec, month: str, org: Any) -> Path:
    return OUTPUT_DIR / (
        f"{month}_{safe_filename_part(org.code)}_{safe_filename_part(org.name)}_{spec.title}.xlsx"
    )


def _workbook_text_and_metadata(path: Path) -> tuple[str, dict[str, str], int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        text_values: list[str] = []
        metadata: dict[str, str] = {}
        labels = {
            "税款所属月份": "month",
            "所属月份": "month",
            "机构代码": "org_code",
            "纳税人识别号": "org_code",
            "机构名称": "org_name",
            "单位名称": "org_name",
        }
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 80), max_col=min(sheet.max_column, 30), values_only=True):
                cells = [str(value).strip() if value is not None else "" for value in row]
                text_values.extend(value for value in cells if value)
                for index, value in enumerate(cells[:-1]):
                    normalized = value.replace(" ", "").rstrip("：:")
                    if normalized in labels and cells[index + 1]:
                        key = labels[normalized]
                        candidate = cells[index + 1].strip()
                        if key == "org_code" and not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{2,49}", candidate):
                            continue
                        if key == "org_name" and not any(
                            token in candidate for token in ("公司", "营业部", "分公司", "事务所", "中心", "单位")
                        ):
                            continue
                        metadata.setdefault(key, candidate)
        return "\n".join(text_values), metadata, len(workbook.sheetnames)
    finally:
        workbook.close()


def validate_report_xlsx(path: Path, spec: ReportSpec, month: str, org_code: str, org_name: str) -> tuple[bool, str]:
    if not path.is_file() or path.stat().st_size <= 1024:
        return False, "文件不存在或小于等于 1 KB"
    if not zipfile.is_zipfile(path):
        return False, "文件不是有效 XLSX/ZIP"
    try:
        text, metadata, sheet_count = _workbook_text_and_metadata(path)
    except Exception as exc:
        return False, f"工作簿无法读取：{exc}"
    if sheet_count < 1:
        return False, "工作簿没有工作表"
    if spec.keyword not in text:
        return False, f"工作簿未包含预期关键字：{spec.keyword}"
    month_value = metadata.get("month", "")
    if month_value:
        normalized = month_value.replace("年", "-").replace("月", "").replace("/", "-")
        found = re.search(r"(\d{4})-(\d{1,2})", normalized)
        if found and f"{found.group(1)}-{int(found.group(2)):02d}" != month:
            return False, f"工作簿月份不匹配：{month_value}"
    metadata_code = metadata.get("org_code", "")
    if metadata_code and org_code not in metadata_code:
        return False, f"工作簿机构代码不匹配：{metadata_code}"
    metadata_name = metadata.get("org_name", "")
    if metadata_name and org_name not in metadata_name and metadata_name not in org_name:
        return False, f"工作簿机构名称不匹配：{metadata_name}"
    return True, "有效"


def quarantine_invalid(path: Path, reason: str) -> Path:
    quarantine = OUTPUT_DIR / "quarantine"
    quarantine.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    target = quarantine / f"{path.stem}_{stamp}{path.suffix}"
    shutil.move(str(path), target)
    log(f"无效历史文件已隔离：{path.name} -> {target.name}；原因：{reason}")
    return target


def resolve_orgs(orgs: list[Any], org_code: str | None, org_name: str | None, start_org_code: str | None) -> list[Any]:
    selected = [org for org in orgs if not org_code or org.code == org_code]
    if org_name:
        selected = [org for org in selected if org.name == org_name or org_name in org.name]
    if start_org_code:
        index = next((i for i, org in enumerate(selected) if org.code == start_org_code), None)
        if index is None:
            raise RuntimeError(f"没有找到开始机构代码：{start_org_code}")
        selected = selected[index:]
    if not selected:
        raise RuntimeError(f"没有找到机构：code={org_code or '-'}，name={org_name or '-'}")
    return selected


def _verify_current_org(page: Any, org: Any, selector: str) -> None:
    value = page.locator(selector).filter(visible=True).first
    value.wait_for(state="visible", timeout=15000)
    text = value.inner_text(timeout=5000)
    if org.code not in text and org.name not in text:
        raise RuntimeError(f"当前机构校验失败：页面显示“{text}”，目标为 {org.name}（{org.code}）")
    log(f"当前机构校验通过：{org.name}（{org.code}）")


def _open_export(page: Any, scope: Any, spec: ReportSpec, config: dict[str, str]) -> None:
    button = scope.locator(config["export_button"]).filter(visible=True).first
    button.wait_for(state="visible", timeout=15000)
    button.click()
    item = page.locator(config["export_menu_item"]).filter(has_text=spec.export_label).filter(visible=True).first
    item.wait_for(state="visible", timeout=10000)
    item.click()
    log(f"已选择导出报表：{spec.export_label}")


def export_record_matches(text: str, org_name: str, month: str, spec: ReportSpec) -> bool:
    compact = re.sub(r"\s+", "", text)
    return all(value in compact for value in (org_name, month.replace("-", ""), spec.keyword, "处理成功", "下载"))


def _download_matching_record(page: Any, org: Any, month: str, spec: ReportSpec, target: Path, config: dict[str, str]) -> None:
    dialog = page.locator(config["export_dialog"]).filter(visible=True).first
    dialog.wait_for(state="visible", timeout=30000)
    month_compact = month.replace("-", "")
    last_status = ""
    for attempt in range(1, 21):
        check_cancelled()
        rows = dialog.locator(config["export_rows"]).filter(has_text=org.name).filter(has_text=month_compact).filter(has_text=spec.keyword)
        for index in range(rows.count()):
            row = rows.nth(index)
            text = row.inner_text(timeout=3000)
            last_status = text[:240].replace("\n", " | ")
            if not export_record_matches(text, org.name, month, spec):
                continue
            download_link = row.locator(config["export_download"]).filter(has_text="下载").filter(visible=True)
            if download_link.count() == 0:
                continue
            with page.expect_download(timeout=30000) as download_info:
                download_link.first.click()
            download_info.value.save_as(str(target))
            log(f"已下载{spec.title}：{target}")
            return
        if attempt < 20:
            log(f"导出记录尚未处理成功，第 {attempt} 次刷新")
            refresh = dialog.locator(config["export_refresh"]).filter(visible=True).first
            refresh.click()
            interruptible_wait(page, 3000)
    raise TimeoutError(f"未找到匹配的成功导出记录：{org.name} / {spec.keyword} / {month_compact}；最后状态：{last_status}")


def _download_one(
    page: Any,
    org: Any,
    month: str,
    spec: ReportSpec,
    config: dict[str, str],
    helpers: dict[str, Any],
) -> bool:
    check_cancelled()
    target = report_path(spec, month, org)
    if target.exists():
        valid, reason = validate_report_xlsx(target, spec, month, org.code, org.name)
        if valid:
            log(f"已有有效文件，跳过重复下载：{target.name}")
            return True
        quarantine_invalid(target, reason)

    page.evaluate(f"location.hash = '{spec.route}'")
    interruptible_wait(page, 1200)
    helpers["close_popups"](page)
    helpers["set_tax_month"](page, month)
    scope = page.locator(config["page_scope"]).filter(visible=True).first
    scope.wait_for(state="visible", timeout=15000)
    status = scope.locator(config["status_value"]).filter(visible=True).first
    status.wait_for(state="visible", timeout=15000)
    status_text = status.inner_text(timeout=5000).strip()
    successful = scope.locator(config["success_record"]).filter(visible=True).first

    def open_export() -> None:
        successful.wait_for(state="visible", timeout=10000)
        _open_export(page, successful, spec, config)

    def download_export() -> Path:
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        _download_matching_record(page, org, month, spec, target, config)
        helpers["close_export_record"](page)
        valid, reason = validate_report_xlsx(target, spec, month, org.code, org.name)
        if not valid:
            quarantine_invalid(target, reason)
            raise RuntimeError(f"下载的 {spec.title} 校验失败：{reason}")
        return target

    downloaded = run_report_download(
        page,
        report_title=f"{org.name}（{org.code}）{spec.title}",
        status_text=status_text,
        open_export=open_export,
        confirm_export=lambda: helpers["confirm_export"](page),
        download_export=download_export,
        success_check=lambda: helpers["export_ready"](page),
        log=log,
    )
    if downloaded is not None:
        log(f"已下载{spec.title}：{downloaded}")
    return downloaded is not None


def capture_failure_evidence(page: Any, org: Any, exc: Exception) -> None:
    try:
        evidence = capture_sanitized_failure(page, OUTPUT_DIR, exc)
        log(f"失败证据已保存：{evidence}")
    except Exception as capture_error:
        log(f"失败证据采集未完成：{capture_error}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="自然人电子税务局分类所得及限售股所得申报表下载")
    parser.add_argument("--cdp", default="http://127.0.0.1:9222")
    parser.add_argument("--month", required=True)
    parser.add_argument("--org-excel", required=True)
    parser.add_argument("--org-code")
    parser.add_argument("--org-name")
    parser.add_argument("--start-org-code")
    parser.add_argument("--yes", action="store_true")
    parser.add_argument("--keep-tabs", action="store_true")
    parser.add_argument("--slow-mo", type=int, default=120)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", args.month):
        raise ValueError("月份必须使用 YYYY-MM 格式")
    config = load_site_config()

    from etax_batch_export import (  # type: ignore[import-not-found]
        close_duplicate_etax_tabs,
        close_common_popups,
        close_export_record_dialog,
        ensure_withholding_page,
        make_context,
        read_orgs_from_excel,
        set_tax_month,
        switch_org,
    )
    from etax_tax_certificate_download import (  # type: ignore[import-not-found]
        confirm_report_export,
        has_visible_confirm_export,
    )
    from playwright.sync_api import sync_playwright
    from etax_popup_guard import set_popup_guard_context

    orgs = resolve_orgs(read_orgs_from_excel(Path(args.org_excel).resolve()), args.org_code, args.org_name, args.start_org_code)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    helpers = {
        "close_popups": close_common_popups,
        "close_export_record": close_export_record_dialog,
        "set_tax_month": set_tax_month,
        "confirm_export": confirm_report_export,
        "export_ready": has_visible_confirm_export,
    }
    log(f"目标税款所属月份：{args.month}；待处理机构数：{len(orgs)}")
    with sync_playwright() as playwright:
        browser, context, page = make_context(playwright, args)
        page.set_default_timeout(15000)
        page.set_default_navigation_timeout(30000)
        if args.cdp and not args.keep_tabs:
            page = close_duplicate_etax_tabs(context, page)
        try:
            page = ensure_withholding_page(page)
            for org in orgs:
                check_cancelled()
                set_popup_guard_context(org_code=org.code, month=args.month, task="extra_income_reports")
                log(f"开始下载扩展申报表：{org.name}（{org.code}）")
                try:
                    switch_org(page, org)
                    page = ensure_withholding_page(page)
                    _verify_current_org(page, org, config["current_org_scope"])
                    count = 0
                    for spec in REPORT_SPECS:
                        check_cancelled()
                        downloaded = _download_one(page, org, args.month, spec, config, helpers)
                        count += int(downloaded)
                        status = "success" if downloaded else "skipped"
                        reason = "downloaded" if downloaded else "no_records"
                        log(f"[RPA_RESULT] org_code={org.code} category={spec.key} status={status} count={int(downloaded)} reason={reason}")
                    log(f"机构扩展申报表下载完成：{org.name}（{org.code}），共 {count} 个文件")
                except Exception as exc:
                    capture_failure_evidence(page, org, exc)
                    raise
        finally:
            if not args.cdp and browser:
                browser.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
