from pathlib import Path
import json
import sqlite3

from openpyxl import load_workbook


ROOT = Path(__file__).parents[1]
REFERENCE = ROOT / "个税测试" / "核对底稿" / "个税核对底稿 - 20260827-有数据.xlsx"


def color(value):
    if value is None:
        return None
    if value.type == "rgb":
        return value.rgb
    if value.type == "indexed":
        return f"indexed:{value.indexed}"
    return f"{value.type}:{value.indexed}"


book = load_workbook(REFERENCE, data_only=False)
layout = []
for sheet in book.worksheets:
    rows = []
    for row in range(1, min(sheet.max_row, 8) + 1):
        populated = []
        for col in range(1, min(sheet.max_column, 60) + 1):
            cell = sheet.cell(row, col)
            if cell.value not in (None, ""):
                populated.append({
                    "coordinate": cell.coordinate,
                    "value": str(cell.value),
                    "fill": color(cell.fill.fgColor),
                    "font": color(cell.font.color) if cell.font.color else None,
                    "bold": cell.font.bold,
                    "alignment": cell.alignment.horizontal,
                })
        if populated:
            rows.append(populated)
    layout.append({
        "name": sheet.title,
        "rows": sheet.max_row,
        "columns": sheet.max_column,
        "freeze": sheet.freeze_panes,
        "merged": [str(item) for item in sheet.merged_cells.ranges],
        "header_rows": rows,
    })

(ROOT / ".tmp_reference_inspection" / "reference_layout.json").write_text(
    json.dumps(layout, ensure_ascii=False, indent=2), encoding="utf-8"
)
print(json.dumps([{key: item[key] for key in ("name", "rows", "columns", "freeze", "merged")} for item in layout], ensure_ascii=False))

db = sqlite3.connect(ROOT / "storage" / "tax_accounting.db")
db.row_factory = sqlite3.Row
master = db.execute(
    "select stored_path from personnel_master_artifacts where company_id=2 and period_id=5 and person_type='broker' and scope_type='month'"
).fetchone()
artifact = db.execute(
    "select a.stored_path from artifacts a join jobs j on j.id=a.job_id where j.company_id=2 and j.period_id=5 and j.workflow_code='broker_tax' and j.status='success' and a.artifact_type='broker_tax_result' order by a.created_at desc limit 1"
).fetchone()
for label, row in (("master", master), ("broker_result", artifact)):
    if not row:
        print(label, "missing")
        continue
    path = Path(row["stored_path"])
    frame_book = load_workbook(path, read_only=True, data_only=True)
    sheet = frame_book.worksheets[0]
    values = list(sheet.iter_rows(values_only=True))
    headers = [str(value or "") for value in values[0]] if values else []
    name_index = next((index for index, value in enumerate(headers) if value in {"*姓名", "姓名", "员工姓名"}), None)
    names = [str(row[name_index] or "") for row in values[1:] if name_index is not None and row[name_index] not in (None, "")]
    print(label, path.name, len(values) - 1, names)
    frame_book.close()
